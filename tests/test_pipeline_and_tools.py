import os
import sys
import pytest
import docx
from docx.oxml.ns import qn

SCRIPTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts"))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from chunk_audio import compute_chunk_ranges
from annotate_docx import (
    style_p,
    add_run,
    extract_page_number,
    annotate_docx
)

def test_compute_chunk_ranges_basic():
    # 1800s (30 min) audio, chunk_len=540, overlap=10
    ranges = compute_chunk_ranges(1800, chunk_len=540, overlap=10)
    assert len(ranges) == 4
    
    # Check first chunk
    assert ranges[0]["index"] == 1
    assert ranges[0]["start_sec"] == 0.0
    assert ranges[0]["duration_sec"] == 550.0  # 540 + 10 overlap
    assert ranges[0]["overlap_sec"] == 0
    
    # Check second chunk
    assert ranges[1]["index"] == 2
    assert ranges[1]["start_sec"] == 540.0
    assert ranges[1]["duration_sec"] == 550.0
    assert ranges[1]["overlap_sec"] == 10

def test_compute_chunk_ranges_ghost_chunk_prevention():
    # Audio with remaining trailing duration <= overlap should not spawn ghost chunk
    ranges = compute_chunk_ranges(545, chunk_len=540, overlap=10)
    # Remaining duration after first chunk starts at 0 is 545. Next chunk would start at 540, remaining 5s <= 10s overlap
    assert len(ranges) == 1
    assert ranges[0]["duration_sec"] == 545.0

def test_style_p_idempotent_xml_tags():
    doc = docx.Document()
    p = doc.add_paragraph("Test Paragraph")
    
    # Call style_p twice
    style_p(p)
    style_p(p)
    
    pPr = p._p.get_or_add_pPr()
    bidi_nodes = pPr.findall(qn('w:bidi'))
    jc_nodes = pPr.findall(qn('w:jc'))
    
    # Exactly one bidi and one jc tag should exist
    assert len(bidi_nodes) == 1
    assert len(jc_nodes) == 1

def test_add_run_font_xml_escaping():
    doc = docx.Document()
    p = doc.add_paragraph()
    
    # Font name containing quotes or angle brackets (XML injection test)
    malicious_font = 'Dubai" <script>evil</script>'
    run = add_run(p, "Sample Text", font_name=malicious_font)
    
    # Check run font name
    assert run.font.name == malicious_font
    # If XML parsing didn't throw and run was successfully added, escaping succeeded
    assert len(p.runs) == 1

def test_extract_page_number_variants():
    assert extract_page_number("Page_1") == 1
    assert extract_page_number("Page 12") == 12
    assert extract_page_number("اسلاید مرتبط ۳") == 3
    assert extract_page_number("اسلاید ۴۵") == 45
    assert extract_page_number("صفحه ۶") == 6
    assert extract_page_number("متن عمومی بدون شماره") is None
    assert extract_page_number("") is None

def test_annotate_docx_input_validation(tmp_path):
    missing_docx = str(tmp_path / "missing.docx")
    missing_json = str(tmp_path / "missing.json")
    out_docx = str(tmp_path / "out.docx")
    
    with pytest.raises(FileNotFoundError, match="Input docx not found"):
        annotate_docx(missing_docx, missing_json, out_docx)
        
    # Create empty docx
    real_docx = str(tmp_path / "real.docx")
    doc = docx.Document()
    doc.save(real_docx)
    
    with pytest.raises(FileNotFoundError, match="Annotations JSON not found"):
        annotate_docx(real_docx, missing_json, out_docx)

def test_annotate_docx_marker_matching(tmp_path):
    import json
    doc_path = str(tmp_path / "test_doc.docx")
    json_path = str(tmp_path / "annotations.json")
    out_path = str(tmp_path / "annotated.docx")
    
    # Document with a regular heading, a numbered slide page, and a body paragraph
    doc = docx.Document()
    doc.add_paragraph("Introduction Heading")
    doc.add_paragraph("Slide Page 1")
    doc.add_paragraph("Unrelated Body Paragraph without numbers")
    doc.save(doc_path)
    
    # Annotations: one numeric marker (Page 1) and one non-numeric heading marker ("Introduction Heading")
    # Also an unmatched marker ("NonExistent Marker") that should NOT falsely match paragraphs with None p_num
    annotations_data = {
        "boxes": [
            {
                "before_marker": "اسلاید ۱",
                "header": "Box 1 - Numeric Match",
                "sections": [{"title": "Note 1", "items": ["Item A"]}]
            },
            {
                "before_marker": "Introduction Heading",
                "header": "Box 2 - Non-Numeric Match",
                "sections": [{"title": "Note 2", "items": ["Item B"]}]
            },
            {
                "before_marker": "NonExistent Marker",
                "header": "Box 3 - Should Not Match",
                "sections": [{"title": "Note 3", "items": ["Item C"]}]
            }
        ]
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(annotations_data, f, ensure_ascii=False)
        
    annotate_docx(doc_path, json_path, out_path)
    
    # Verify the output document has exactly 2 callout tables inserted (Box 1 and Box 2), NOT 3
    annotated_doc = docx.Document(out_path)
    assert len(annotated_doc.tables) == 2

def test_build_pamphlet_json_validation(tmp_path):
    from create_slide_pamphlet import build_pamphlet_from_json
    
    missing_file = str(tmp_path / "does_not_exist.json")
    out_docx = str(tmp_path / "out.docx")
    
    # 1. Non-existent file
    with pytest.raises(FileNotFoundError, match="not found"):
        build_pamphlet_from_json(missing_file, out_docx)
        
    # 2. Malformed JSON
    corrupt_json = tmp_path / "corrupt.json"
    corrupt_json.write_text("{ unclosed json", encoding="utf-8")
    with pytest.raises(ValueError, match="Malformed JSON"):
        build_pamphlet_from_json(str(corrupt_json), out_docx)
        
    # 3. Non-list JSON root
    dict_json = tmp_path / "dict.json"
    dict_json.write_text('{"slides": []}', encoding="utf-8")
    with pytest.raises(ValueError, match="Expected a JSON list"):
        build_pamphlet_from_json(str(dict_json), out_docx)

def test_add_cross_reference_box_execution():
    from create_slide_pamphlet import add_cross_reference_box
    
    doc = docx.Document()
    # Test forward reference
    add_cross_reference_box(doc, target_slide=5, note="مراجعه به اسلاید ۵ جهت مقایسه", direction="forward")
    # Test backward reference
    add_cross_reference_box(doc, target_slide=2, note="همان‌طور که در اسلاید ۲ دیدیم", direction="backward")
    
    # Both calls should succeed and create table boxes
    assert len(doc.tables) == 2


def test_transcribe_chunks_mime_type_and_policy():
    from transcribe_chunks import get_mime_type
    assert get_mime_type("chunk_01.m4a") == "audio/mp4"
    assert get_mime_type("audio.mp3") == "audio/mp3"
    assert get_mime_type("speech.wav") == "audio/wav"
    assert get_mime_type("recording.ogg") == "audio/ogg"


def test_create_slide_pamphlet_advisory_banner_on_missing_chart(tmp_path):
    from create_slide_pamphlet import build_pamphlet_from_json
    import json
    
    slides_data = [
        {
            "slide_number": 39,
            "title_fa": "نمودار مقایسه شیب بروز عوارض",
            "title_en": "Microvascular Risk Curve",
            "has_charts": True,
            "bullets": [
                {"lead": "نکته کلیدی", "text": "همبستگی مستقیم HbA1c و عوارض"}
            ]
        }
    ]
    json_path = tmp_path / "translated.json"
    json_path.write_text(json.dumps(slides_data, ensure_ascii=False), encoding="utf-8")
    out_docx = str(tmp_path / "pamphlet.docx")
    
    build_pamphlet_from_json(str(json_path), out_docx, img_dir=str(tmp_path / "empty_imgs"))
    assert os.path.exists(out_docx)
    
    doc = docx.Document(out_docx)
    text_content = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                text_content += "\n" + "\n".join(p.text for p in cell.paragraphs)
                
    assert "تذکر آموزشی: این اسلاید حاوی نمودار/شکل تخصصی است" in text_content


def test_build_pamphlet_inherits_visual_flags_and_image_from_raw_slides(tmp_path):
    """Verifies that build_pamphlet_from_json inherits has_images and img_path from raw_slides.json."""
    from create_slide_pamphlet import build_pamphlet_from_json
    import json
    
    # Create fake image
    img_dir = tmp_path / "slide_images"
    img_dir.mkdir()
    fake_img = str(img_dir / "slide_33.png")
    with open(fake_img, "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\nfake")
        
    raw_slides = [
        {
            "slide_number": 33,
            "title_raw": "Charcot Neuroarthropathy Clinical Presentation",
            "has_images": True,
            "img_path": fake_img
        }
    ]
    raw_path = tmp_path / "raw_slides.json"
    raw_path.write_text(json.dumps(raw_slides, ensure_ascii=False), encoding="utf-8")
    
    # translated_slides.json has NO has_images or img_path (simulating agent omitting them)
    trans_slides = [
        {
            "slide_number": 33,
            "title_fa": "تظاهرات بالینی نوروآرتروپاتی شارکو",
            "title_en": "Charcot Neuroarthropathy",
            "bullets": [
                {"lead": "مشخصات بالینی", "text": "تغییر شکل استخوانی و زخم مفاصل پا"}
            ]
        }
    ]
    trans_path = tmp_path / "translated.json"
    trans_path.write_text(json.dumps(trans_slides, ensure_ascii=False), encoding="utf-8")
    
    out_docx = str(tmp_path / "pamphlet_inherited.docx")
    build_pamphlet_from_json(str(trans_path), out_docx, img_dir=str(img_dir), raw_slides_path=str(raw_path))
    assert os.path.exists(out_docx)
    
    doc = docx.Document(out_docx)
    # Check that doc has an embedded image picture
    has_picture = False
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        if "drawing" in run._r.xml:
                            has_picture = True
    assert has_picture is True, "Expected image to be embedded in docx via raw_slides metadata inheritance"


def test_build_pamphlet_renders_table_structural_advisory_on_missing_table_data(tmp_path):
    """Verifies that when raw_slides has has_tables: True but translation only has bullets, a warning is rendered."""
    from create_slide_pamphlet import build_pamphlet_from_json
    import json
    
    raw_slides = [
        {
            "slide_number": 17,
            "title_raw": "Statin Therapy Intensity Classification",
            "has_tables": True
        }
    ]
    raw_path = tmp_path / "raw_slides.json"
    raw_path.write_text(json.dumps(raw_slides, ensure_ascii=False), encoding="utf-8")
    
    trans_slides = [
        {
            "slide_number": 17,
            "title_fa": "طبقه‌بندی شدت درمان با استاتین",
            "title_en": "Statin Therapy Intensity Classification",
            "bullets": [
                {"lead": "دوز بالا", "text": "آتورواستاتین ۴۰ الی ۸۰ میلی‌گرم"}
            ]
            # No table_data provided
        }
    ]
    trans_path = tmp_path / "translated.json"
    trans_path.write_text(json.dumps(trans_slides, ensure_ascii=False), encoding="utf-8")
    
    out_docx = str(tmp_path / "pamphlet_table_warning.docx")
    build_pamphlet_from_json(str(trans_path), out_docx, raw_slides_path=str(raw_path))
    assert os.path.exists(out_docx)
    
    doc = docx.Document(out_docx)
    all_text = ""
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                all_text += "\n" + "\n".join(p.text for p in c.paragraphs)
    assert "تذکر ساختاری: اسلاید اصلی حاوی جدول است" in all_text


def test_run_pipeline_args_forwarding(monkeypatch, tmp_path):
    """
    Regression Test (Phase 1 Feedback Loop):
    Verifies that run_pipeline.py properly forwards --img-dir to verify_slide_alignment.py
    and --raw to create_slide_pamphlet.py.
    """
    import run_pipeline
    import argparse

    commands_executed = []
    def fake_run_cmd(cmd_list, description=""):
        commands_executed.append(cmd_list)
        return True

    monkeypatch.setattr(run_pipeline, "run_cmd", fake_run_cmd)

    # 1. Test cmd_verify_slides forwarding --img-dir
    raw_f = str(tmp_path / "raw.json")
    trans_f = str(tmp_path / "trans.json")
    with open(raw_f, "w") as f: f.write("[]")
    with open(trans_f, "w") as f: f.write("[]")

    args_v = argparse.Namespace(
        raw=raw_f,
        translated=trans_f,
        auto_fix=False,
        ref_corpus=None,
        allow_review=False,
        img_dir="./custom_images"
    )
    run_pipeline.cmd_verify_slides(args_v)
    last_cmd = commands_executed[-1]
    assert "--img-dir" in last_cmd, f"Expected --img-dir to be forwarded in verify_slide_alignment command, got: {last_cmd}"
    assert "./custom_images" in last_cmd

    # 2. Test cmd_build forwarding --raw
    args_b = argparse.Namespace(
        translated=trans_f,
        output="out.docx",
        img_dir="./custom_images",
        title="Custom Title",
        ref_book="هاریسون",
        verify=False,
        raw=raw_f
    )
    run_pipeline.cmd_build(args_b)
    last_build_cmd = commands_executed[-1]
    assert "--raw" in last_build_cmd, f"Expected --raw to be forwarded in create_slide_pamphlet command, got: {last_build_cmd}"
    assert raw_f in last_build_cmd


def test_render_reference_table_list_rows_and_length_mismatch():
    """
    Regression Test (Phase 1 Feedback Loop):
    Verifies that render_reference_table in create_slide_pamphlet.py gracefully supports
    standard 2D list rows: [["val1", "val2"], ...], handles column count mismatches without
    IndexError, and does not crash with AttributeError when rows are lists instead of dicts.
    """
    from create_slide_pamphlet import render_reference_table
    import docx

    doc = docx.Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)

    headers = ["شاخص بالینی", "مقدار هدف"]
    # Test combination of:
    # 1. Standard 2D list row
    # 2. Row with extra columns (length mismatch)
    # 3. Row with fewer columns
    # 4. Dictionary row with 'cols'
    # 5. Dictionary banner row
    items = [
        ["HbA1c", "< 7.0%"],
        ["فشار خون سیستولیک", "< 130 mmHg", "ستون اضافی که نباید کرش کند"],
        ["کلسترول LDL"],
        {"type": "banner", "text": "اهداف دارودرمانی خط اول"},
        {"cols": ["آتورواستاتین", "۲۰ الی ۴۰ میلی‌گرم"]}
    ]

    # Prior to fix, this crashed with AttributeError: 'list' object has no attribute 'get'
    render_reference_table(cell, headers, items)
    
    # Assert table was added inside cell
    assert len(cell.tables) == 1
    t = cell.tables[0]
    assert len(t.rows) == len(items) + 1  # 1 header row + 5 item rows
    assert "HbA1c" in t.rows[1].cells[0].text
    assert "7.0" in t.rows[1].cells[1].text
    assert "اهداف دارودرمانی" in t.rows[4].cells[0].text
    assert "آتورواستاتین" in t.rows[5].cells[0].text


def test_transcribe_chunks_retry_backoff(monkeypatch, tmp_path):
    """
    Regression Test (Phase 1 Feedback Loop):
    Verifies that transcribe_chunk_with_gemini_rest in transcribe_chunks.py
    implements automatic exponential backoff retry on HTTP 429 rate limit errors
    and returns successfully when a retry succeeds.
    """
    import transcribe_chunks
    import urllib.error
    from unittest.mock import MagicMock
    import io
    import json

    chunk_file = tmp_path / "test_chunk.m4a"
    chunk_file.write_bytes(b"dummy audio binary data")

    call_count = 0
    def fake_urlopen(req, timeout=180):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            # First attempt: simulate HTTP 429 Too Many Requests
            fp = io.BytesIO(b'{"error": {"code": 429, "message": "Resource has been exhausted"}}')
            raise urllib.error.HTTPError(
                url="https://generativelanguage.googleapis.com",
                code=429,
                msg="Too Many Requests",
                hdrs={},
                fp=fp
            )
        else:
            # Second attempt: succeed with transcript JSON
            mock_resp = MagicMock()
            resp_data = {
                "candidates": [
                    {
                        "content": {
                            "parts": [
                                {"text": "تدریس بالینی بیماری کوشینگ و هیپرکورتیزولیسم"}
                            ]
                        }
                    }
                ]
            }
            mock_resp.read.return_value = json.dumps(resp_data).encode("utf-8")
            mock_resp.__enter__.return_value = mock_resp
            return mock_resp

    monkeypatch.setattr("urllib.request.urlopen", fake_urlopen)
    # Monkeypatch time.sleep to avoid waiting in tests
    monkeypatch.setattr("time.sleep", lambda s: None)

    res = transcribe_chunks.transcribe_chunk_with_gemini_rest(
        str(chunk_file),
        api_key="fake-test-key",
        max_retries=3,
        retry_delay=0.01
    )

    assert call_count == 2, f"Expected exactly 2 attempts, but urlopen was called {call_count} times"
    assert "تدریس بالینی بیماری کوشینگ" in res
