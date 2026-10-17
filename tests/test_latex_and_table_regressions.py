import os
import sys
import json
import pytest
import docx

SCRIPTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts"))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from text_utils import sanitize_latex_math, normalize_for_display
from create_slide_pamphlet import add_bidi_text, add_r
from verify_slide_alignment import check_slide_pair


def test_sanitize_latex_math_with_spaces_in_text_wrappers_and_subscripts():
    """
    Regression Test (Axis 1: LaTeX Math Leaks):
    Model outputs '$1, 25 (\\ text { OH }) _2 \\ text {D}$'.
    sanitize_latex_math must clean whitespace after backslash in wrappers and convert _2 to ₂.
    """
    raw = r"$1, 25 (\ text { OH }) _2 \ text {D}$"
    cleaned = sanitize_latex_math(raw)
    
    # Must not contain backslashes or raw latex tokens
    assert "\\" not in cleaned, f"Found stray backslash in '{cleaned}'"
    assert "text" not in cleaned, f"Found stray 'text' in '{cleaned}'"
    assert "{" not in cleaned and "}" not in cleaned, f"Found stray braces in '{cleaned}'"
    assert "$" not in cleaned, f"Found stray dollar signs in '{cleaned}'"
    assert "_2" not in cleaned, f"Subscript _2 was not converted in '{cleaned}'"
    assert "₂" in cleaned, f"Expected subscript ₂ in '{cleaned}'"
    assert "(OH)₂D" in cleaned.replace(" ", ""), f"Expected (OH)₂D in '{cleaned}'"


def test_sanitize_latex_math_single_subscripts_and_spaces():
    """
    Regression Test:
    Tests single-char subscripts without braces like H_2O, CO_2, and optional spaces around _ and ^.
    """
    cases = [
        (r"H_2O و CO_2", "H₂O و CO₂"),
        (r"T_3 و T_4", "T₃ و T₄"),
        (r"V_1 تا V_6", "V₁ تا V₆"),
        (r"Ca^{2+} و HCO_3^-", "Ca²⁺ و HCO₃⁻"),
        (r"\ text{Na}^+ و \  mathrm{K}^+", "Na⁺ و K⁺"),
        (r"A _ 1 و B ^ 2", "A₁ و B²")
    ]
    for raw, expected in cases:
        result = sanitize_latex_math(raw)
        assert result == expected, f"Failed for '{raw}': got '{result}', expected '{expected}'"


def test_bidi_run_renders_explicit_ltr_val_zero():
    """
    Regression Test (Axis 1: BiDi Isolation in Word OpenXML):
    add_r with is_rtl=False must explicitly emit <w:rtl w:val="0"/> in rPr,
    not simply remove the element, so Word never inherits paragraph-level RTL bidi.
    """
    doc = docx.Document()
    p = doc.add_paragraph()
    run = add_r(p, "1,25(OH)₂D", is_rtl=False)
    
    rPr = run._r.get_or_add_rPr()
    rtl_elem = rPr.find(docx.oxml.ns.qn('w:rtl'))
    assert rtl_elem is not None, "Expected <w:rtl> element to be present for explicit LTR isolation"
    val = rtl_elem.get(docx.oxml.ns.qn('w:val'))
    assert val == "0" or val == "false", f"Expected w:val='0' or 'false' for LTR run, got '{val}'"


def test_add_bidi_text_keeps_subscripts_and_formulas_in_ltr():
    """
    Regression Test:
    In 'سطح 1,25(OH)₂D در سرم', 1,25(OH)₂D must be rendered as a single LTR run,
    and subscripts ₀-₉ must not be fragmented into adjacent RTL runs.
    """
    doc = docx.Document()
    p = doc.add_paragraph()
    add_bidi_text(p, "سطح 1,25(OH)₂D در سرم")
    
    # Check runs in paragraph
    ltr_runs = [r.text.strip() for r in p.runs if r._r.find(docx.oxml.ns.qn('w:rPr')) is not None and 
                r._r.find(docx.oxml.ns.qn('w:rPr')).find(docx.oxml.ns.qn('w:rtl')) is not None and
                r._r.find(docx.oxml.ns.qn('w:rPr')).find(docx.oxml.ns.qn('w:rtl')).get(docx.oxml.ns.qn('w:val')) == "0"]
    
    assert any("1,25(OH)₂D" in r or "(OH)₂D" in r for r in ltr_runs), (
        f"Expected 1,25(OH)₂D in LTR runs, got: {ltr_runs} across all runs: {[r.text for r in p.runs]}"
    )


def test_image_table_without_ocr_tokens_does_not_silently_pass_empty_stub():
    """
    Regression Test (Axis 2: Image-based Table Bypass):
    Slide has has_image_table: True, but ocr_text_lines is empty ([]).
    If translated slide contains only 1 stub row, check_slide_pair must NOT silently pass;
    it must flag INCOMPLETE_TABLE_TRANSCRIPTION or MISSING_TABLE_DATA.
    """
    raw_slide = {
        "slide_number": 45,
        "title_raw": "Classification of Causes of Hypercalcemia",
        "text_lines": ["Classification of Causes of Hypercalcemia"],
        "ocr_text_lines": [],  # Empty OCR!
        "has_images": True,
        "has_tables": True,
        "has_image_table": True
    }

    # Deficient stub translation with only 1 row
    trans_slide = {
        "slide_number": 45,
        "title_fa": "طبقه‌بندی علل هایپرکلسمی",
        "title_en": "Classification of Causes of Hypercalcemia",
        "bullets": [],
        "table_data": {
            "headers": ["دسته علل"],
            "rows": [
                {"cols": ["علل متفرقه"]}
            ]
        }
    }

    issues = check_slide_pair(raw_slide, trans_slide)
    err_types = [i.get("type") or i.get("error_type") for i in issues]
    assert any(k in err_types for k in ("INCOMPLETE_TABLE_TRANSCRIPTION", "MISSING_TABLE_DATA", "INSUFFICIENT_TABLE_DATA")), (
        f"Expected table incompleteness error for stub image table without OCR, got issues: {issues}"
    )


def test_table_recall_threshold_blocks_half_omitted_clinical_tables():
    """
    Regression Test (Axis 2: Table Recall Strictness):
    When an agent translates only 50% of a reference classification table (e.g. omitting Renal Failure & Vitamin D),
    the table completeness gate must fail (threshold >= 0.70).
    """
    raw_slide_45 = {
        "slide_number": 45,
        "title_raw": "Classification of Causes of Hypercalcemia",
        "text_lines": ["Classification of Causes of Hypercalcemia"],
        "ocr_text_lines": [
            "I. Parathyroid-related: Primary hyperparathyroidism, Lithium therapy, Familial hypocalciuric hypercalcemia",
            "II. Malignancy-related: Solid tumor metastases breast, lung, kidney; Hematologic multiple myeloma, lymphoma, leukemia",
            "III. Vitamin D-related: Vitamin D intoxication, 1,25(OH)2D sarcoidosis",
            "IV. Associated with Renal Failure: Severe secondary hyperparathyroidism, Aluminum intoxication, Milk-alkali syndrome"
        ],
        "has_images": True,
        "has_tables": True,
        "has_image_table": True
    }

    # Deficient translation that omits half the table
    trans_half_omitted = {
        "slide_number": 45,
        "title_fa": "طبقه‌بندی علل هایپرکلسمی",
        "title_en": "Classification of Causes of Hypercalcemia",
        "bullets": [],
        "table_data": {
            "headers": ["دسته علل", "اتیولوژی‌های اختصاصی"],
            "rows": [
                {"cols": ["وابسته به پاراتیروئید", "هایپرپاراتیروئیدیسم اولیه و درمان با لیتیوم"]},
                {"cols": ["وابسته به بدخیمی", "تومورهای توپر پستان، ریه، کلیه"]}
            ]
        }
    }

    issues = check_slide_pair(raw_slide_45, trans_half_omitted)
    incomplete_errs = [i for i in issues if i.get("type") == "INCOMPLETE_TABLE_TRANSCRIPTION"]
    assert len(incomplete_errs) >= 1, f"Expected INCOMPLETE_TABLE_TRANSCRIPTION for half-omitted table, got: {issues}"
    assert incomplete_errs[0]["severity"] == "error"
