import os
import sys
import pytest

SCRIPTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts"))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from text_utils import normalize_persian_text, extract_tokens

def test_persian_arabic_char_normalization():
    # Arabic kaf and ya
    raw = "كتاب پزشكي و غدد"
    expected = "کتاب پزشکی و غدد"
    assert normalize_persian_text(raw) == expected

    # Ta marbuta
    assert normalize_persian_text("غدة") == "غده"

    # Hamza variants
    assert normalize_persian_text("أنسولین") == "انسولین"
    assert normalize_persian_text("إندوکرین") == "اندوکرین"
    assert normalize_persian_text("آدرنال") == "ادرنال"
    assert normalize_persian_text("مؤثر") == "موثر"
    assert normalize_persian_text("مسئله") == "مسیله"

def test_digit_normalization():
    # Persian digits
    p_text = "اسلاید شماره ۱۲۳۴۵ و ۶۷۸۹۰"
    assert "12345" in normalize_persian_text(p_text)
    assert "67890" in normalize_persian_text(p_text)

    # Arabic digits
    a_text = "چانک ٠١٢٣٤ و ٥٦٧٨٩"
    assert "01234" in normalize_persian_text(a_text)
    assert "56789" in normalize_persian_text(a_text)

def test_tashkeel_and_zero_width_space():
    # Diacritics (harakat/tashkeel)
    tashkeel_text = "مُقَدِّمَةٌ فیزیولوژی"
    assert normalize_persian_text(tashkeel_text) == "مقدمه فیزیولوژی"

    # Zero-width non-joiner
    zwnj_text = "بیماری\u200cهای هورمونی\u200b"
    assert normalize_persian_text(zwnj_text) == "بیماری های هورمونی"

def test_token_extraction_medical_acronyms():
    lines = ["هورمون GH و گیرنده D2 و همچنین مسیر IP3 و غلظت کلسیم خون"]
    tokens = extract_tokens(lines)
    
    # English clinical abbreviations
    assert "gh" in tokens
    assert "d2" in tokens
    assert "ip3" in tokens
    
    # Persian substantive clinical terms
    assert "هورمون" in tokens
    assert "گیرنده" in tokens
    assert "کلسیم" in tokens
    
    # Common stopwords should be excluded
    assert "این" not in tokens
    assert "و" not in tokens

def test_persian_punctuation_and_conversational_stopwords():
    # Verify Persian punctuation (، ؛ ؟) separation prevents word clinging
    punct_text = "آدرنال، هیپوفیز؛ تیروئید؟"
    norm_text = normalize_persian_text(punct_text)
    assert "ادرنال" in norm_text
    assert "هیپوفیز" in norm_text
    assert "تیرویید" in norm_text
    assert "،" not in norm_text
    assert "؛" not in norm_text
    assert "؟" not in norm_text

    tokens = extract_tokens(punct_text)
    assert "ادرنال" in tokens
    assert "هیپوفیز" in tokens
    assert "تیرویید" in tokens

    # Verify conversational stopwords & verbal auxiliaries are stripped
    colloquial_sentence = "ببینید این آزمایش رو باید بکنیم که مشخص بشه و معلوم باشه که بیمار چشه"
    colloq_tokens = extract_tokens(colloquial_sentence)
    assert "ازمایش" in colloq_tokens
    assert "بیمار" in colloq_tokens
    assert "بکنیم" not in colloq_tokens
    assert "بشه" not in colloq_tokens
    assert "باشه" not in colloq_tokens
    assert "ببینید" not in colloq_tokens

def test_display_vs_matching_normalization():
    from text_utils import normalize_for_display, normalize_for_matching
    
    # Correct Persian orthography must be preserved in display
    raw = "مسئله آدرنال و غده هیپوفیز"
    disp = normalize_for_display(raw)
    assert "مسئله" in disp
    assert "آدرنال" in disp
    
    # Matching normalizes for fuzzy lexical lookup
    match = normalize_for_matching(raw)
    assert "مسیله" in match
    assert "ادرنال" in match
    
    # Display must preserve scientific arrows and chemical notation
    eq = "CO2 + H2O ⇌ H2CO3 → H+ + HCO3-"
    disp_eq = normalize_for_display(eq)
    assert "⇌" in disp_eq
    assert "→" in disp_eq
    assert "HCO3-" in disp_eq

def test_medical_tokens_compounds_and_dosages():
    text = "بیمار با دوز 25mg داروی مسدودکننده β1 تحت درمان است. همچنین IGFBP-3 و Ca²⁺ و HCO₃⁻ و TNF-α و T3/T4 و 7% شیوع بررسی شد."
    tokens = extract_tokens(text)
    
    # Hyphenated/slashed compounds
    assert "igfbp-3" in tokens
    assert "tnf-α" in tokens
    assert "t3/t4" in tokens
    
    # Chemical / ion notations with sub/superscript
    assert "ca²⁺" in tokens
    assert "hco₃⁻" in tokens
    assert "β1" in tokens
    
    # Dosages and percentages
    assert "25mg" in tokens
    assert "7%" in tokens

def test_clean_markdown_text():
    from text_utils import clean_markdown_text

    raw = """### قطعه اول: مبحث تیروئید
#### بیانات استاد:
استاد فرمودند **بیماری کوشینگ** مهم است.
- **نکته کلیدی:** سطح ACTH بالا است.
* درمان: جراحی."""

    cleaned = clean_markdown_text(raw)
    assert "*" not in cleaned
    assert "#" not in cleaned
    assert "\n" not in cleaned
    assert "بیانات استاد" not in cleaned
    assert "بیماری کوشینگ مهم است" in cleaned
    assert "سطح ACTH بالا است" in cleaned

def test_parse_inline_spans():
    from text_utils import parse_inline_spans

    text = "استاد فرمودند **بیماری کوشینگ (Cushing)** همراه با *استریا* و دوز 25mg است."
    runs = parse_inline_spans(text)

    # Must have extracted runs with appropriate bold/italic
    bold_runs = [r for r in runs if r["bold"]]
    italic_runs = [r for r in runs if r["italic"]]
    plain_runs = [r for r in runs if not r["bold"] and not r["italic"]]

    assert len(bold_runs) == 1
    assert "بیماری کوشینگ (Cushing)" in bold_runs[0]["text"]
    assert len(italic_runs) == 1
    assert "استریا" in italic_runs[0]["text"]
    assert any("25mg" in r["text"] for r in plain_runs)

    # Zero asterisks in all extracted text
    assert all("*" not in r["text"] for r in runs)
    assert all("\n" not in r["text"] for r in runs)

def test_parse_markdown_blocks_reflows_single_newlines():
    from text_utils import parse_markdown_blocks

    text = """### مبحث تنظیم هورمونی
#### بیانات استاد:
استاد در تشریح این مبحث فرمودند
که محور هیپوتالاموس-هیپوفیز
نقش تنظیمی اصلی را بر عهده دارد.

- **نکته مهم:** فیدبک منفی اعمال می‌شود.
- آزمایش تشخیصی: سطح هورمون سنجیده می‌شود.

1. مرحله اول آزمایش
2. مرحله دوم تفسیر نتایج

در پایان استاد به موارد منع مصرف اشاره کردند."""

    blocks = parse_markdown_blocks(text)
    
    # 1. Generic header filtered
    assert not any("بیانات استاد" in b.get("text", "") for b in blocks)

    # 2. Heading parsed
    head_blocks = [b for b in blocks if b["type"] == "heading"]
    assert len(head_blocks) == 1
    assert head_blocks[0]["text"] == "مبحث تنظیم هورمونی"
    assert "#" not in head_blocks[0]["text"]

    # 3. Consecutive lines reflowed into single continuous paragraph without \n
    para_blocks = [b for b in blocks if b["type"] == "paragraph"]
    assert len(para_blocks) == 2
    assert "\n" not in para_blocks[0]["text"]
    assert "استاد در تشریح این مبحث فرمودند که محور هیپوتالاموس-هیپوفیز نقش تنظیمی اصلی را بر عهده دارد." in para_blocks[0]["text"]

    # 4. Bullets parsed
    bullet_blocks = [b for b in blocks if b["type"] == "bullet"]
    assert len(bullet_blocks) == 2
    assert bullet_blocks[0]["lead"] == "نکته مهم:"
    assert "فیدبک منفی اعمال می‌شود." in bullet_blocks[0]["text"]
    assert bullet_blocks[1]["lead"] == "آزمایش تشخیصی:"

    # 5. Numbered list parsed
    num_blocks = [b for b in blocks if b["type"] == "numbered"]
    assert len(num_blocks) == 2
    assert num_blocks[0]["num"] == "1"
    assert num_blocks[1]["num"] == "2"

def test_create_slide_pamphlet_renders_zero_newlines_in_runs_and_justified(tmp_path):
    import json
    import docx
    from create_slide_pamphlet import build_pamphlet_from_json

    slides = [
        {
            "slide_number": 1,
            "title_fa": "فیزیولوژی غدد",
            "title_en": "Endocrine Physiology",
            "spoken_lecture": """### مبحث اختصاصی
#### بیانات استاد:
استاد فرمودند **بیماری کوشینگ** مهم است
و باید به دقت بررسی گردد
چرا که علائم گوناگون دارد.

- **نکته کلیدی:** سطح کورتیزول سنجیده می‌شود.
- آزمایش تکمیلی: تست دگزامتازون.""",
            "bullets": [
                {"lead": "نکته اصلی:", "text": "تنظیم هورمونی **بسیار مهم** است\nو باید رعایت شود."},
                ("تست آزمایشگاهی:", "آزمایش خون ناشتا")
            ],
            "ref_note": "شرح تکمیلی رفرنس هاریسون جهت تفهیم مبحث به صورت **کامل** و دقیق."
        }
    ]

    json_path = tmp_path / "test_slides.json"
    json_path.write_text(json.dumps(slides, ensure_ascii=False), encoding="utf-8")
    out_docx = tmp_path / "test_out.docx"

    build_pamphlet_from_json(str(json_path), str(out_docx))

    doc = docx.Document(str(out_docx))
    
    # 1. Assert ZERO literal \n or \r characters inside ANY run
    for p in doc.paragraphs:
        for r in p.runs:
            assert "\n" not in r.text, f"Found newline in run: {r.text!r}"
            assert "\r" not in r.text, f"Found carriage return in run: {r.text!r}"
            assert "**" not in r.text, f"Found raw markdown asterisks in run: {r.text!r}"
            assert "###" not in r.text, f"Found raw markdown hashes in run: {r.text!r}"
            assert "####" not in r.text, f"Found raw markdown hashes in run: {r.text!r}"

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        assert "\n" not in r.text, f"Found newline in table run: {r.text!r}"
                        assert "\r" not in r.text, f"Found carriage return in table run: {r.text!r}"
                        assert "**" not in r.text, f"Found raw markdown asterisks in table run: {r.text!r}"

    # 2. Assert paragraphs that have text have justification enabled
    justified_count = 0
    for p in doc.paragraphs:
        if p.text.strip() and not p.text.startswith("🎙️") and not p.text.startswith("جزوه"):
            if p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY:
                justified_count += 1
    assert justified_count > 0, "No justified paragraphs found"


def test_sanitize_ref_note_removes_duplicate_prefixes():
    from text_utils import sanitize_ref_note, has_ref_note_prefix

    cases = [
        (
            "💡 شرح تکمیلی رفرنس (هاریسون) جهت تفهیم مبحث: قشر غده فوق‌کلیوی از سه رده هورمونی تشکیل شده است.",
            "قشر غده فوق‌کلیوی از سه رده هورمونی تشکیل شده است."
        ),
        (
            "شرح تکمیلی رفرنس (هاریسون): قشر غده فوق‌کلیوی از سه رده هورمونی تشکیل شده است.",
            "قشر غده فوق‌کلیوی از سه رده هورمونی تشکیل شده است."
        ),
        (
            "💡 شرح تکمیلی رفرنس (هاریسون) جهت تفهیم مبحث: 💡 شرح تکمیلی رفرنس (هاریسون) جهت تفهیم مبحث: هورمون‌های استروئیدی سنتز می‌شوند.",
            "هورمون‌های استروئیدی سنتز می‌شوند."
        ),
        (
            "توضیحات تکمیلی (Cecil): بررسی بالینی نشان‌دهنده ترشح آلدوسترون است.",
            "بررسی بالینی نشان‌دهنده ترشح آلدوسترون است."
        ),
        (
            "نکات تکمیلی رفرنس - هورمون آلدوسترون در تنظیم سدیم نقش دارد.",
            "هورمون آلدوسترون در تنظیم سدیم نقش دارد."
        ),
        (
            "💡 قشر غده فوق‌کلیوی هورمون‌های استروئیدی ترشح می‌کند.",
            "قشر غده فوق‌کلیوی هورمون‌های استروئیدی ترشح می‌کند."
        ),
        (
            "Reference (Harrison): Adrenal cortex secretes three classes of steroid hormones.",
            "Adrenal cortex secretes three classes of steroid hormones."
        ),
        (
            "💡 Ref Note: Adrenal cortex glucocorticoids regulate metabolism.",
            "Adrenal cortex glucocorticoids regulate metabolism."
        )
    ]

    for raw, expected in cases:
        assert has_ref_note_prefix(raw) is True
        assert sanitize_ref_note(raw) == expected


def test_sanitize_ref_note_preserves_clean_text():
    from text_utils import sanitize_ref_note, has_ref_note_prefix

    clean_cases = [
        "قشر غده فوق‌کلیوی از سه رده هورمونی گلوکوکورتیکوئیدها تشکیل شده است.",
        "شرح حال بیمار نشان‌دهنده علائم افزایش کورتیزول و سندرم کوشینگ بود.",
        "رفرنس اصلی این مبحث کتاب هاریسون است و باید مطالعه شود.",
        "هورمون آلدوسترون تنظیم‌کننده هومئوستاز سدیم و پتاسیم است."
    ]

    for text in clean_cases:
        assert has_ref_note_prefix(text) is False
        assert sanitize_ref_note(text) == text

