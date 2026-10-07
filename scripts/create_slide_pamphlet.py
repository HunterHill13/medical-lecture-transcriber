#!/usr/bin/env python3
"""
create_slide_pamphlet.py: Builds a modern, publication-grade Persian medical
study guide Word document (.docx) from extracted presentation slides (PPTX/PDF)
and class audio lecture transcription.

Features:
- Unified Dubai font across all text runs
- Explicit Complex Script XML font sizing (w:szCs) & bolding (w:bCs) for Microsoft Word
- Balanced heading hierarchy (Title 16pt, TOC 14pt, H1 14pt, H2 11pt)
- Soft light sky blue slide boxes (#85C1E9 border, #EBF5FB header, #FFFFFF body)
- Programmatic JSON slide integration (zero hallucination)
- Reference-accurate Harrison table rendering (peach headers & merged slate-blue banners)
- Bidirectional text sanitization & English parenthetical run isolation (zero flipping)
- Supplementary reference commentary for skipped slides (ref_note)
- Student review badges for textless/image-only slides
- Single-placement figure embedding (no duplication)
- Classroom Q&A callouts
"""

import os
import sys
import json
import argparse
import re
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from version import VERSION

from text_utils import (
    is_ceremonial_slide,
    clean_markdown_text,
    parse_inline_spans,
    parse_markdown_blocks,
    flatten_spoken_lecture
)

def convert_image_to_png(src_path: str, dst_path: str = None) -> str:
    """
    Converts vector, metafile, and specialized image formats (WMF, EMF, SVG, WebP, TIFF, BMP)
    to standard high-resolution PNG.
    Returns path to converted PNG if successful, or original path/None on failure.
    """
    if not src_path or not os.path.exists(src_path):
        return None
    base, ext = os.path.splitext(src_path)
    ext_l = ext.lower()
    if ext_l == ".png":
        return src_path
        
    if not dst_path:
        dst_path = base + ".png"
        
    # 1. Try Pillow (supports WMF via WmfImagePlugin, WebP, TIFF, BMP)
    try:
        from PIL import Image, WmfImagePlugin
        with Image.open(src_path) as img:
            img.convert("RGB").save(dst_path, "PNG")
            if os.path.exists(dst_path):
                return dst_path
    except Exception:
        pass

    # 2. Try PyMuPDF (supports SVG, PDF, and various image formats)
    try:
        import pymupdf
        doc = pymupdf.open(src_path)
        if len(doc) > 0:
            pix = doc[0].get_pixmap(dpi=150)
            pix.save(dst_path)
            if os.path.exists(dst_path):
                return dst_path
    except Exception:
        pass

    # 3. Windows Native fallback: System.Drawing via PowerShell (native support for WMF, EMF, BMP, etc.)
    if sys.platform == "win32":
        try:
            import subprocess
            ps_script = (
                f"Add-Type -AssemblyName System.Drawing; "
                f"$img = [System.Drawing.Image]::FromFile('{os.path.abspath(src_path)}'); "
                f"$img.Save('{os.path.abspath(dst_path)}', [System.Drawing.Imaging.ImageFormat]::Png); "
                f"$img.Dispose();"
            )
            subprocess.run(
                ["powershell", "-NoProfile", "-Command", ps_script],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                check=True,
                timeout=15
            )
            if os.path.exists(dst_path):
                return dst_path
        except Exception:
            pass

    return src_path

def find_and_prepare_slide_image(img_dir: str, slide_num: int) -> str | None:
    """
    Dynamically searches img_dir for any candidate image representing slide_num,
    converting vector/metafile images (WMF/EMF/SVG) to standard PNG if needed.
    """
    if not img_dir or not os.path.isdir(img_dir):
        return None
        
    exts = ("png", "jpg", "jpeg", "wmf", "emf", "webp", "tiff", "tif", "bmp", "svg")
    patterns = [
        f"slide_{slide_num:03d}",
        f"slide_{slide_num:02d}",
        f"slide_{slide_num}",
        f"slide_{slide_num:03d}_img",
        f"slide_{slide_num:02d}_img",
        f"slide_{slide_num}_img"
    ]
    
    # 1. Exact priority matches
    for pat in patterns:
        for ext in exts:
            cand = os.path.join(img_dir, f"{pat}.{ext}")
            if os.path.exists(cand):
                return convert_image_to_png(cand)
                
    # 2. Dynamic directory scan matching slide number prefix
    try:
        files = sorted(os.listdir(img_dir))
        target_prefixes = (f"slide_{slide_num:02d}", f"slide_{slide_num:03d}", f"slide_{slide_num}_", f"slide_{slide_num}.")
        for f in files:
            f_lower = f.lower()
            if any(f_lower.startswith(p) for p in target_prefixes):
                if any(f_lower.endswith(f".{ext}") for ext in exts):
                    cand = os.path.join(img_dir, f)
                    return convert_image_to_png(cand)
    except Exception:
        pass
        
    return None


def set_p_rtl(p, space_before=2, space_after=3, line_spacing=1.15, align_justify=False):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if align_justify:
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p._p.get_or_add_pPr()
    if pPr.find(docx.oxml.ns.qn('w:bidi')) is None:
        pPr.append(parse_xml(r'<w:bidi {}/>'.format(nsdecls('w'))))

def add_r(p, text, font_name="Dubai", size_pt=11, bold=False, italic=False, color_rgb=(0x26, 0x26, 0x26), is_rtl=True):
    if text is not None:
        # Sanitize any raw newline characters from text runs to avoid corrupting Word's justification
        text = str(text).replace('\r\n', ' ').replace('\n', ' ').replace('\r', ' ')
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    
    rFonts = rPr.find(docx.oxml.ns.qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(r'<w:rFonts {} w:ascii="{}" w:hAnsi="{}" w:cs="{}"/>'.format(nsdecls('w'), font_name, font_name, font_name))
        rPr.append(rFonts)
    else:
        rFonts.set(docx.oxml.ns.qn('w:cs'), font_name)
        rFonts.set(docx.oxml.ns.qn('w:ascii'), font_name)
        rFonts.set(docx.oxml.ns.qn('w:hAnsi'), font_name)
        
    rtl_elem = rPr.find(docx.oxml.ns.qn('w:rtl'))
    if is_rtl:
        if rtl_elem is None:
            rPr.append(parse_xml(r'<w:rtl {} w:val="1"/>'.format(nsdecls('w'))))
        else:
            rtl_elem.set(docx.oxml.ns.qn('w:val'), '1')
    else:
        if rtl_elem is not None:
            rPr.remove(rtl_elem)
            
    # Set explicit Complex Script font size (w:szCs) matching w:sz (in half-points)
    # Microsoft Word uses w:szCs for Persian/Arabic/Complex Script text.
    sz_half_pts = str(int(round(size_pt * 2)))
    
    sz_elem = rPr.find(docx.oxml.ns.qn('w:sz'))
    if sz_elem is None:
        rPr.append(parse_xml(r'<w:sz {} w:val="{}"/>'.format(nsdecls('w'), sz_half_pts)))
    else:
        sz_elem.set(docx.oxml.ns.qn('w:val'), sz_half_pts)
        
    szCs_elem = rPr.find(docx.oxml.ns.qn('w:szCs'))
    if szCs_elem is None:
        rPr.append(parse_xml(r'<w:szCs {} w:val="{}"/>'.format(nsdecls('w'), sz_half_pts)))
    else:
        szCs_elem.set(docx.oxml.ns.qn('w:val'), sz_half_pts)
        
    # Set explicit Complex Script bold (w:bCs) and italic (w:iCs)
    bCs_elem = rPr.find(docx.oxml.ns.qn('w:bCs'))
    if bold:
        if bCs_elem is None:
            rPr.append(parse_xml(r'<w:bCs {}/>'.format(nsdecls('w'))))
    else:
        if bCs_elem is not None:
            rPr.remove(bCs_elem)
            
    iCs_elem = rPr.find(docx.oxml.ns.qn('w:iCs'))
    if italic:
        if iCs_elem is None:
            rPr.append(parse_xml(r'<w:iCs {}/>'.format(nsdecls('w'))))
    else:
        if iCs_elem is not None:
            rPr.remove(iCs_elem)
            
    return run

def add_bidi_text(p, text, font_name="Dubai", size_pt=10.5, bold=False, italic=False, color_rgb=(0x26, 0x26, 0x26)):
    """
    Renders mixed Persian/English scientific text with strict BiDi isolation.
    Preserves arrows (→, ⇌, <->, ->), chemical formulas (Ca²⁺, HCO₃⁻), inequalities (p < 0.05),
    and Latin abbreviations without converting scientific symbols to Persian prose.
    """
    if not text:
        return
    text = str(text).replace('\r\n', ' ').replace('\n', ' ').replace('\r', ' ')
    pattern = r'(\([A-Za-z0-9_\-\s,\./%α-ωΑ-Ω→⇌⇄<>=\+\^±]+\)|[A-Za-z0-9_\-\./%α-ωΑ-Ω\+\^±]+(?:\s*(?:→|->|⇌|<->|⇄|<=|>=|<|>|=)\s*[A-Za-z0-9_\-\./%α-ωΑ-Ω\+\^±]+)+|[A-Za-z0-9_\-\./%α-ωΑ-Ω\+\^±]{2,}|[→⇌⇄]|(?:<=|>=|[<>=])\s*\d+(?:\.\d+)?)'
    tokens = re.split(pattern, text)
    for tok in tokens:
        if not tok:
            continue
        has_persian = any('\u0600' <= c <= '\u06FF' for c in tok)
        if not has_persian and (re.search(r'[A-Za-z0-9→⇌⇄<>=]', tok) or tok.startswith('(')):
            add_r(p, " " + tok.strip() + " ", font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb, is_rtl=False)
        else:
            add_r(p, tok, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb, is_rtl=True)

def add_formatted_bidi_text(p, text_or_runs, font_name="Dubai", size_pt=10.5, default_bold=False, default_italic=False, color_rgb=(0x26, 0x26, 0x26)):
    """
    Renders text with inline markdown bold (**...**) and italic (*...*),
    stripping raw markdown artifacts and ensuring clean BiDi isolation and smooth Word justification.
    """
    if not text_or_runs:
        return
    if isinstance(text_or_runs, list) and text_or_runs and isinstance(text_or_runs[0], dict):
        runs = text_or_runs
    else:
        runs = parse_inline_spans(str(text_or_runs), default_bold=default_bold, default_italic=default_italic)
        
    for r in runs:
        r_text = r.get("text", "")
        if not r_text:
            continue
        r_bold = r.get("bold", default_bold)
        r_italic = r.get("italic", default_italic)
        add_bidi_text(p, r_text, font_name=font_name, size_pt=size_pt, bold=r_bold, italic=r_italic, color_rgb=color_rgb)

def add_numbered_p(doc, num, body):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=2, space_after=2, align_justify=True)
    add_r(p, f"{num}. ", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    add_formatted_bidi_text(p, body, font_name="Dubai", size_pt=11, color_rgb=(0x26, 0x26, 0x26))
    return p

def add_title(doc, main_title, sub_title=""):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=16, space_after=4)
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    add_r(p, main_title, font_name="Dubai", size_pt=16, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    if sub_title:
        p2 = doc.add_paragraph()
        set_p_rtl(p2, space_before=2, space_after=14)
        p2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        add_r(p2, sub_title, font_name="Dubai", size_pt=11, italic=True, color_rgb=(0x4A, 0x55, 0x68))

def add_toc(doc, rows_data):
    p_toc_title = doc.add_paragraph()
    set_p_rtl(p_toc_title, space_before=14, space_after=8)
    add_r(p_toc_title, "📑 فهرست عناوین، زمان فایل صوتی و اسلایدهای مرتبط", font_name="Dubai", size_pt=14, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    
    headers = ["بخش", "موضوع و مبحث تدریس کلاسی استاد", "زمان فایل صوتی", "اسلایدهای مرتبط"]
    t = doc.add_table(rows=len(rows_data) + 1, cols=4)
    tPr = t._tbl.tblPr
    tPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    for c_idx, h in enumerate(headers):
        c = t.cell(0, c_idx)
        tcPr = c._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd {} w:fill="78281F"/>'.format(nsdecls('w'))))
        p = c.paragraphs[0]
        set_p_rtl(p, space_before=3.5, space_after=3.5)
        add_r(p, h, font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0xFF, 0xFF, 0xFF))
        
    for r_idx, rdata in enumerate(rows_data):
        bg = "FAF7F5" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(rdata):
            c = t.cell(r_idx + 1, c_idx)
            tcPr = c._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg)))
            p = c.paragraphs[0]
            set_p_rtl(p, space_before=2.5, space_after=2.5)
            # Section: Burgundy (#78281F), Audio time: Deep Orange (#D35400), Others: Charcoal (#262626)
            c_rgb = (0x78, 0x28, 0x1F) if c_idx == 0 else ((0xD3, 0x54, 0x00) if c_idx == 2 else (0x26, 0x26, 0x26))
            add_r(p, val, font_name="Dubai", size_pt=10, bold=(c_idx == 0 or c_idx == 2), color_rgb=c_rgb)
            
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(12)

add_toc_table = add_toc

def add_h1(doc, text, audio_time=None):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=18, space_after=5)
    add_r(p, text, font_name="Dubai", size_pt=14, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    if audio_time:
        add_r(p, f"   [⏱️ زمان فایل صوتی: {audio_time}]", font_name="Dubai", size_pt=10, bold=True, color_rgb=(0xD3, 0x54, 0x00))
    return p

def add_h2(doc, text, audio_time=None):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=11, space_after=3)
    add_r(p, text, font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x0E, 0x62, 0x51))
    if audio_time:
        add_r(p, f"   [⏱️ {audio_time}]", font_name="Dubai", size_pt=9.5, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    return p

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=2.5, space_after=3.5, align_justify=True)
    if bold_prefix:
        clean_prefix = clean_markdown_text(bold_prefix).strip()
        add_r(p, clean_prefix + " ", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    add_formatted_bidi_text(p, text, font_name="Dubai", size_pt=11, color_rgb=(0x26, 0x26, 0x26))
    return p

def add_bullet_p(doc, lead, body):
    p = doc.add_paragraph()
    set_p_rtl(p, space_before=2, space_after=2, align_justify=True)
    clean_lead = clean_markdown_text(lead).strip() if lead else ""
    if clean_lead:
        if not clean_lead.endswith(":"):
            clean_lead += ":"
        add_r(p, "• " + clean_lead + " ", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x0E, 0x62, 0x51))
    else:
        add_r(p, "• ", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x0E, 0x62, 0x51))
    add_formatted_bidi_text(p, body, font_name="Dubai", size_pt=11, color_rgb=(0x26, 0x26, 0x26))
    return p

def add_figure_with_caption(doc, img_path, caption_title, caption_text="", max_width_in=5.2):
    if img_path and os.path.exists(img_path):
        final_img = convert_image_to_png(img_path)
        if final_img and os.path.exists(final_img):
            p = doc.add_paragraph()
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(final_img, width=Inches(max_width_in))
            
            cp = doc.add_paragraph()
            set_p_rtl(cp, space_before=2, space_after=8)
            cp.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            add_r(cp, f"🖼️ {caption_title} ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x78, 0x28, 0x1F))
            if caption_text:
                add_r(cp, f"— {caption_text}", font_name="Dubai", size_pt=10, italic=True, color_rgb=(0x55, 0x55, 0x55))
            return True
    return False

def add_qa_box(doc, question, answer):
    tbl = doc.add_table(rows=1, cols=1)
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="FBF9FD"/>'.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="6" w:space="0" w:color="E8DAEF"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="E8DAEF"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="E8DAEF"/>
            <w:right w:val="single" w:sz="36" w:space="0" w:color="6C3483"/>
        </w:tcBorders>
    '''.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    
    p0 = cell.paragraphs[0]
    set_p_rtl(p0, space_before=2, space_after=2)
    add_r(p0, "❓ پرسش دانشجو در کلاس: ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x6C, 0x34, 0x83))
    add_r(p0, question, font_name="Dubai", size_pt=10.5, italic=True, color_rgb=(0x33, 0x33, 0x33))
    
    p1 = cell.add_paragraph()
    set_p_rtl(p1, space_before=2, space_after=2)
    add_r(p1, "💡 پاسخ استاد: ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1E, 0x84, 0x49))
    add_bidi_text(p1, answer, font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_student_review_badge(doc, slide_num, note_text=""):
    tbl = doc.add_table(rows=1, cols=1)
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="FEF9E7"/>'.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="8" w:space="0" w:color="FAD7A0"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="FAD7A0"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="FAD7A0"/>
            <w:right w:val="single" w:sz="36" w:space="0" w:color="D35400"/>
        </w:tcBorders>
    '''.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    
    p = cell.paragraphs[0]
    set_p_rtl(p, space_before=2, space_after=2)
    add_r(p, f"⚠️ برچسب بازبینی دانشجو (اسلاید {slide_num} - فاقد متن صریح در پاورپوینت): ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0xD3, 0x54, 0x00))
    desc = note_text or "این اسلاید در فایل پاورپوینت فاقد عنوان یا متن تشریحی بود؛ با توجه به توالی اسلایدهای قبل و بعد و مبحث تدریس استاد، در این بخش از جزوه جانمایی شد."
    add_r(p, desc, font_name="Dubai", size_pt=10, italic=True, color_rgb=(0x55, 0x55, 0x55))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_cross_reference_box(doc, target_slide, note="", direction="forward"):
    """Renders an elegant blue callout box representing the professor's cross-reference to another slide."""
    tbl = doc.add_table(rows=1, cols=1)
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="F4F6F7"/>'.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="6" w:space="0" w:color="D5DBDB"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="D5DBDB"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D5DBDB"/>
            <w:right w:val="single" w:sz="36" w:space="0" w:color="2471A3"/>
        </w:tcBorders>
    '''.format(nsdecls('w'))))
    tcPr.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="100" w:type="dxa"/>
            <w:bottom w:w="100" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="180" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    
    p = cell.paragraphs[0]
    set_p_rtl(p, space_before=2, space_after=2)
    dir_icon = "⏩" if direction == "forward" else "⏪"
    target_str = f" به اسلاید {target_slide}" if target_slide else " به مباحث دیگر"
    add_r(p, f"🔗 ارجاع استاد در تدریس{target_str} {dir_icon}: ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x24, 0x71, 0xA3))
    if note:
        add_bidi_text(p, note, font_name="Dubai", size_pt=10, italic=True, color_rgb=(0x33, 0x33, 0x33))
        
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def render_reference_table(parent_cell, headers, items, ref_book_name="هاریسون"):
    """Renders a medical reference-styled table with peach headers and merged category banners."""
    tp = parent_cell.add_paragraph()
    set_p_rtl(tp, space_before=4, space_after=3)
    add_r(tp, f"📊 جدول تفصیلی ترجمه‌شده فارسی (منطبق دقیق بر ساختار و رفرنس {ref_book_name}):", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    
    t = parent_cell.add_table(rows=len(items) + 1, cols=len(headers))
    tPr = t._tbl.tblPr
    tPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    # Header Row (Peach / Cream background like Harrison original: #F5EBE1, Text #5D4037)
    for c_idx, h in enumerate(headers):
        c = t.cell(0, c_idx)
        tcPr_t = c._tc.get_or_add_tcPr()
        tcPr_t.append(parse_xml(r'<w:shd {} w:fill="F5EBE1"/>'.format(nsdecls('w'))))
        tcPr_t.append(parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="single" w:sz="8" w:space="0" w:color="5D4037"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/>
                <w:bottom w:val="single" w:sz="8" w:space="0" w:color="5D4037"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/>
            </w:tcBorders>
        '''.format(nsdecls('w'))))
        tcPr_t.append(parse_xml(r'''
            <w:tcMar {} >
                <w:top w:w="90" w:type="dxa"/>
                <w:bottom w:w="90" w:type="dxa"/>
                <w:left w:w="110" w:type="dxa"/>
                <w:right w:w="110" w:type="dxa"/>
            </w:tcMar>
        '''.format(nsdecls('w'))))
        p_t = c.paragraphs[0]
        set_p_rtl(p_t, space_before=2.5, space_after=2.5)
        add_r(p_t, h, font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x5D, 0x40, 0x37))
        
    for r_idx, item in enumerate(items):
        row_num = r_idx + 1
        if item.get("type") == "banner":
            cell_start = t.cell(row_num, 0)
            cell_end = t.cell(row_num, len(headers) - 1)
            cell_merged = cell_start.merge(cell_end)
            
            tcPr_t = cell_merged._tc.get_or_add_tcPr()
            tcPr_t.append(parse_xml(r'<w:shd {} w:fill="4A709C"/>'.format(nsdecls('w'))))
            tcPr_t.append(parse_xml(r'''
                <w:tcBorders {} >
                    <w:top w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
                    <w:left w:val="none"/>
                    <w:bottom w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            '''.format(nsdecls('w'))))
            tcPr_t.append(parse_xml(r'''
                <w:tcMar {} >
                    <w:top w:w="80" w:type="dxa"/>
                    <w:bottom w:w="80" w:type="dxa"/>
                    <w:left w:w="120" w:type="dxa"/>
                    <w:right w:w="120" w:type="dxa"/>
                </w:tcMar>
            '''.format(nsdecls('w'))))
            p_m = cell_merged.paragraphs[0]
            set_p_rtl(p_m, space_before=2.5, space_after=2.5)
            add_r(p_m, item["text"], font_name="Dubai", size_pt=11, bold=True, color_rgb=(0xFF, 0xFF, 0xFF))
        else:
            cols_val = item["cols"]
            bg = item.get("bg", "FFFFFF" if r_idx % 2 == 0 else "FDFBF7")
            for c_idx, val in enumerate(cols_val):
                c = t.cell(row_num, c_idx)
                tcPr_t = c._tc.get_or_add_tcPr()
                tcPr_t.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg)))
                tcPr_t.append(parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                    </w:tcBorders>
                '''.format(nsdecls('w'))))
                tcPr_t.append(parse_xml(r'''
                    <w:tcMar {} >
                        <w:top w:w="70" w:type="dxa"/>
                        <w:bottom w:w="70" w:type="dxa"/>
                        <w:left w:w="90" w:type="dxa"/>
                        <w:right w:w="90" w:type="dxa"/>
                    </w:tcMar>
                '''.format(nsdecls('w'))))
                p_c = c.paragraphs[0]
                set_p_rtl(p_c, space_before=2, space_after=2)
                is_bold = (c_idx == 0) and not val.startswith("   ")
                add_bidi_text(p_c, val, font_name="Dubai", size_pt=10, bold=is_bold, color_rgb=(0x26, 0x26, 0x26))

render_harrison_table = render_reference_table

def add_slide_box_from_json(doc, slide_data, img_path=None, table_data=None, ref_book_name="کتاب مرجع", image_already_shown=False):
    """Renders a 2-row slide box strictly from JSON data with soft sky blue border and white body.
    Supports pure text slides, localized tables, and optional diagrams.
    Auto-detects table_data from slide_data if not explicitly passed."""
    slide_num = slide_data.get('slide_number', '?')
    title_fa = slide_data.get('title_fa', f'اسلاید {slide_num}')
    title_en = slide_data.get('title_en', '')
    bullets = slide_data.get('bullets', [])
    ref_note = slide_data.get('ref_note', '')
    is_visual = bool(
        slide_data.get("has_table") or 
        slide_data.get("has_tables") or 
        slide_data.get("has_diagram") or 
        slide_data.get("is_image_only") or 
        slide_data.get("has_figure") or 
        slide_data.get("has_images") or 
        slide_data.get("has_charts") or 
        slide_data.get("has_smartart") or 
        slide_data.get("has_image_table") or 
        slide_data.get("is_visual") or 
        bool(slide_data.get("visual_content_type"))
    )
    is_ceremonial = bool(slide_data.get("is_ceremonial") or is_ceremonial_slide(slide_data))
    is_skipped = bool(slide_data.get("is_skipped") or slide_data.get("unvoiced") or slide_data.get("skipped_by_professor"))
    
    # Auto-fallback to slide_data table_data if omitted from arguments
    if table_data is None:
        table_data = slide_data.get("table_data")
    
    tbl = doc.add_table(rows=2, cols=1)
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    # Row 0: Slim Header Bar (Soft Light Sky Blue #EBF5FB with Soft Sky Blue Border #85C1E9)
    cell0 = tbl.cell(0, 0)
    tcPr0 = cell0._tc.get_or_add_tcPr()
    tcPr0.append(parse_xml(r'<w:shd {} w:fill="EBF5FB"/>'.format(nsdecls('w'))))
    tcPr0.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D4E6F1"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
        </w:tcBorders>
    '''.format(nsdecls('w'))))
    tcPr0.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="80" w:type="dxa"/>
            <w:bottom w:w="80" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="160" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    
    p0 = cell0.paragraphs[0]
    set_p_rtl(p0, space_before=1.5, space_after=2)
    prefix_label = "📑 اسلاید مقدماتی " if is_ceremonial else "📑 اسلاید مرتبط "
    add_r(p0, f"{prefix_label}{slide_num}: ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
    add_r(p0, title_fa, font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x78, 0x28, 0x1F))
    if title_en:
        add_r(p0, f" ({title_en})", font_name="Dubai", size_pt=9.5, italic=True, color_rgb=(0x55, 0x55, 0x55), is_rtl=False)
    if is_ceremonial:
        add_r(p0, " [اسلاید تشریفاتی / مقدماتی]", font_name="Dubai", size_pt=9.5, bold=True, color_rgb=(0x2E, 0x86, 0xC1))
    elif is_skipped:
        add_r(p0, " [⚠️ تدریس‌نشده در کلاس]", font_name="Dubai", size_pt=9.5, bold=True, color_rgb=(0xD3, 0x54, 0x00))
    else:
        audio_range = slide_data.get("audio_time_range") or slide_data.get("audio_time")
        if audio_range:
            add_r(p0, f"  [⏱️ زمان تدریس: {audio_range}]", font_name="Dubai", size_pt=9.5, bold=True, color_rgb=(0xD3, 0x54, 0x00))
        
    # Row 1: Content Area (White background, 1pt Soft Sky Blue Border #85C1E9)
    cell1 = tbl.cell(1, 0)
    tcPr1 = cell1._tc.get_or_add_tcPr()
    tcPr1.append(parse_xml(r'<w:shd {} w:fill="FFFFFF"/>'.format(nsdecls('w'))))
    tcPr1.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="85C1E9"/>
        </w:tcBorders>
    '''.format(nsdecls('w'))))
    tcPr1.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="100" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    
    first_item = True
    # Embed screenshot for visual slides (diagram, table, or image-only) unless already placed in Track 1
    embedded_successfully = False
    if is_visual and img_path and os.path.exists(img_path) and not image_already_shown:
        try:
            final_img = convert_image_to_png(img_path)
            if final_img and os.path.exists(final_img):
                ip = cell1.paragraphs[0] if first_item else cell1.add_paragraph()
                first_item = False
                ip.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                ip.paragraph_format.space_before = Pt(4)
                ip.paragraph_format.space_after = Pt(6)
                run = ip.add_run()
                run.add_picture(final_img, width=Inches(4.8))
                embedded_successfully = True
        except Exception as e:
            sys.stderr.write(f"WARNING: Failed to embed picture for slide {slide_num}: {e}\n")

    if not embedded_successfully and not image_already_shown:
        if slide_data.get("image_extraction_error"):
            ip = cell1.paragraphs[0] if first_item else cell1.add_paragraph()
            first_item = False
            set_p_rtl(ip, space_before=2, space_after=3)
            add_r(ip, "⚠️ [هشدار سیستم: استخراج تصویر این اسلاید با خطا مواجه شد؛ لطفاً فایل اصلی اسلاید را بررسی کنید.]",
                  font_name="Dubai", size_pt=9.5, italic=True, color_rgb=(0xC0, 0x39, 0x2B))
        elif is_visual and (slide_data.get("has_charts") or slide_data.get("has_images") or slide_data.get("is_pure_visual")):
            ip = cell1.paragraphs[0] if first_item else cell1.add_paragraph()
            first_item = False
            set_p_rtl(ip, space_before=2, space_after=3)
            add_r(ip, "⚠️ [تذکر آموزشی: این اسلاید حاوی نمودار/شکل تخصصی است؛ جهت مشاهده تصویر به فایل ارائه اصلی مراجعه نمایید.]",
                  font_name="Dubai", size_pt=9.5, italic=True, color_rgb=(0xD3, 0x54, 0x00))
        
    for b in bullets:
        bp = cell1.paragraphs[0] if first_item else cell1.add_paragraph()
        first_item = False
        set_p_rtl(bp, space_before=1.5, space_after=2, align_justify=True)
        if isinstance(b, dict):
            lead = clean_markdown_text(b.get("lead", "")).strip()
            text = b.get("text", "")
            if lead:
                if not lead.endswith(":"):
                    lead += ":"
                add_r(bp, "• " + lead + " ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
            else:
                add_r(bp, "• ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
            add_formatted_bidi_text(bp, text, font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
        elif isinstance(b, (tuple, list)):
            if len(b) >= 2:
                lead, text = clean_markdown_text(str(b[0])).strip(), " ".join(str(x) for x in b[1:])
            elif len(b) == 1:
                lead, text = clean_markdown_text(str(b[0])).strip(), ""
            else:
                lead, text = "", ""
            if lead:
                if not lead.endswith(":"):
                    lead += ":"
                add_r(bp, "• " + lead + " ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
            else:
                add_r(bp, "• ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
            if text:
                add_formatted_bidi_text(bp, text, font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
        else:
            raw_str = str(b).strip()
            # If string starts with bold-like pattern e.g. "عنوان: متن" or "**عنوان:** متن"
            if ":" in raw_str and not raw_str.startswith("http"):
                parts = raw_str.split(":", 1)
                lead_cand = clean_markdown_text(parts[0]).strip()
                if len(lead_cand.split()) <= 6:
                    add_r(bp, "• " + lead_cand + ": ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
                    add_formatted_bidi_text(bp, parts[1].strip(), font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
                else:
                    add_r(bp, "• ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
                    add_formatted_bidi_text(bp, raw_str, font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
            else:
                add_r(bp, "• ", font_name="Dubai", size_pt=10.5, bold=True, color_rgb=(0x1B, 0x4F, 0x72))
                add_formatted_bidi_text(bp, raw_str, font_name="Dubai", size_pt=10.5, color_rgb=(0x26, 0x26, 0x26))
            
    if is_ceremonial and not bullets and not table_data and not (img_path and os.path.exists(img_path)):
        cp = cell1.paragraphs[0]
        set_p_rtl(cp, space_before=2, space_after=2)
        add_r(cp, title_fa or "اسلاید تشریفاتی / مقدماتی جلسه", font_name="Dubai", size_pt=10.5, italic=True, color_rgb=(0x55, 0x55, 0x55))
        
    if table_data:
        if isinstance(table_data, dict):
            headers = table_data.get("headers", [])
            items_data = table_data.get("rows") or table_data.get("items", [])
        elif isinstance(table_data, (list, tuple)) and len(table_data) == 2:
            headers, items_data = table_data
        else:
            headers, items_data = [], []
        if headers and items_data:
            render_reference_table(cell1, headers, items_data, ref_book_name=ref_book_name)
    elif (slide_data.get("has_tables") or slide_data.get("has_image_table")) and not is_ceremonial:
        tp = cell1.add_paragraph() if not first_item else cell1.paragraphs[0]
        first_item = False
        set_p_rtl(tp, space_before=2, space_after=3)
        add_r(tp, "⚠️ [تذکر ساختاری: اسلاید اصلی حاوی جدول است؛ جهت پایش سلول‌به‌سلول اطلاعات به فایل ارائه مراجعه فرمایید.]",
              font_name="Dubai", size_pt=9.5, italic=True, color_rgb=(0xD3, 0x54, 0x00))
        
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(3)

    if ref_note and not is_ceremonial:
        # Track 3: Separate, dedicated Reference Note Callout Box (OUTSIDE Slide Box table)
        ref_tbl = doc.add_table(rows=1, cols=1)
        ref_tblPr = ref_tbl._tbl.tblPr
        ref_tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
        ref_tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
        
        rcell = ref_tbl.cell(0, 0)
        rcellPr = rcell._tc.get_or_add_tcPr()
        rcellPr.append(parse_xml(r'<w:shd {} w:fill="FBF9FC"/>'.format(nsdecls('w'))))
        rcellPr.append(parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="D7BDE2"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D7BDE2"/>
                <w:right w:val="single" w:sz="18" w:space="0" w:color="6C3483"/>
            </w:tcBorders>
        '''.format(nsdecls('w'))))
        rcellPr.append(parse_xml(r'''
            <w:tcMar {} >
                <w:top w:w="80" w:type="dxa"/>
                <w:bottom w:w="80" w:type="dxa"/>
                <w:left w:w="140" w:type="dxa"/>
                <w:right w:w="160" w:type="dxa"/>
            </w:tcMar>
        '''.format(nsdecls('w'))))
        
        rp = rcell.paragraphs[0]
        set_p_rtl(rp, space_before=2, space_after=2, align_justify=True)
        add_r(rp, f"💡 شرح تکمیلی رفرنس ({ref_book_name}) جهت تفهیم مبحث: ", font_name="Dubai", size_pt=10, bold=True, color_rgb=(0x6C, 0x34, 0x83))
        add_formatted_bidi_text(rp, ref_note, font_name="Dubai", size_pt=10, default_italic=True, color_rgb=(0x33, 0x33, 0x33))

        sp2 = doc.add_paragraph()
        sp2.paragraph_format.space_before = Pt(0)
        sp2.paragraph_format.space_after = Pt(4)

def check_font_availability(font_name="Dubai"):
    """Checks if the requested font is installed on the host OS across Windows, macOS, and Linux."""
    import platform
    import subprocess
    system = platform.system()
    try:
        if system == "Windows":
            import winreg
            reg_paths = [
                r"Software\Microsoft\Windows NT\CurrentVersion\Fonts",
                r"Software\Microsoft\Windows\CurrentVersion\Fonts"
            ]
            for reg_path in reg_paths:
                try:
                    with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, reg_path) as key:
                        for i in range(winreg.QueryInfoKey(key)[1]):
                            name, _, _ = winreg.EnumValue(key, i)
                            if font_name.lower() in name.lower():
                                return True
                except Exception:
                    pass
            # Check C:\Windows\Fonts directly
            windir = os.environ.get("WINDIR", r"C:\Windows")
            fonts_dir = os.path.join(windir, "Fonts")
            if os.path.exists(fonts_dir):
                for f in os.listdir(fonts_dir):
                    if font_name.lower() in f.lower():
                        return True
            return False
        elif system == "Darwin":
            mac_dirs = ["/Library/Fonts", "/System/Library/Fonts", os.path.expanduser("~/Library/Fonts")]
            for d in mac_dirs:
                if os.path.exists(d):
                    for f in os.listdir(d):
                        if font_name.lower() in f.lower():
                            return True
            return False
        else:
            res = subprocess.run(["fc-list", f":family={font_name}"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
            return font_name.lower() in res.stdout.lower()
    except Exception:
        return True

DISCIPLINE_BOOK_MAP = [
    (re.compile(r"غدد|دیابت|کلسیم|پاراتیروئید|هیپوفیز|تیروئید|گوارش|کبد|کلیه|روماتو|خون|هماتو|ریه|نفرول|اندوکرین|انکو|endo|nephro|gastro|rheum", re.I), "هاریسون"),
    (re.compile(r"قلب|عروق|آریتمی|انفارکتوس|نارسایی قلب|کاردیو|ecg|ekg|cardio|valve", re.I), "برانوالد"),
    (re.compile(r"جراحی|آپاندیس|تروما|فتق|سوختگی|لارنگو|surgery|surgical|trauma|hernia", re.I), "شوارتز"),
    (re.compile(r"کودک|اطفال|نوزاد|واکسیناسیون|pediatric|neonat|infant", re.I), "نلسون"),
    (re.compile(r"زنان|زایمان|مامایی|بارداری|سزارین|gynecol|obstet|pregnancy", re.I), "ویلیامز"),
    (re.compile(r"پاتولوژی|آسیب‌شناسی|نئوپلاسم|تومور|بیوپسی|pathol|neoplas", re.I), "رابینز"),
    (re.compile(r"فارماکو|دارو|گیرنده|آنتاگونیست|آگونیست|فارما|pharmaco|drug", re.I), "کاتزونگ"),
    (re.compile(r"فیزیولوژی|پتانسیل عمل|سیناپس|هموستاز|physiol|homeostasis", re.I), "گایتون"),
]

def detect_ref_book(text_corpus, default="هاریسون"):
    if not text_corpus:
        return default
    for pattern, book_name in DISCIPLINE_BOOK_MAP:
        if pattern.search(text_corpus):
            return book_name
    return default

def add_spoken_lecture(doc, spoken_text, audio_time=None):
    """Renders Track 1: spoken classroom lecture narrative associated with a slide/topic,
    with intelligent markdown parsing, header sanitization, bullet recognition, and smooth justification."""
    if not spoken_text:
        return
        
    p_h = doc.add_paragraph()
    set_p_rtl(p_h, space_before=12, space_after=3)
    add_r(p_h, "🎙️ تدریس و بیانات کلاسی استاد", font_name="Dubai", size_pt=11, bold=True, color_rgb=(0x0E, 0x62, 0x51))
    if audio_time:
        add_r(p_h, f"   [⏱️ زمان فایل صوتی: دقیقه {audio_time}]", font_name="Dubai", size_pt=9.5, bold=True, color_rgb=(0xD3, 0x54, 0x00))
        
    blocks = parse_markdown_blocks(spoken_text)
    for block in blocks:
        b_type = block.get("type")
        if b_type == "heading":
            add_h2(doc, block.get("text", ""))
        elif b_type == "bullet":
            add_bullet_p(doc, block.get("lead", ""), block.get("text", ""))
        elif b_type == "numbered":
            add_numbered_p(doc, block.get("num", "1"), block.get("text", ""))
        elif b_type == "paragraph":
            p = doc.add_paragraph()
            set_p_rtl(p, space_before=2.5, space_after=3.5, align_justify=True)
            add_formatted_bidi_text(p, block.get("text", ""), font_name="Dubai", size_pt=11, color_rgb=(0x26, 0x26, 0x26))

def build_pamphlet_from_json(translated_json_path, output_docx_path, title="جزوه جامع پزشکی", ref_book=None, img_dir=None, raw_slides_path=None):
    if not os.path.isfile(translated_json_path):
        raise FileNotFoundError(f"Translated slides JSON file not found: '{translated_json_path}'")
    try:
        with open(translated_json_path, "r", encoding="utf-8") as f:
            slides = json.load(f)
    except json.JSONDecodeError as e:
        raise ValueError(f"Malformed JSON in '{translated_json_path}': {e}") from e

    if not isinstance(slides, list):
        raise ValueError(f"Expected a JSON list of slide objects in '{translated_json_path}', got {type(slides).__name__}")
        
    # Auto-discover raw_slides.json if not explicitly provided
    if not raw_slides_path and os.path.isfile("raw_slides.json"):
        raw_slides_path = "raw_slides.json"
        
    raw_db = {}
    if raw_slides_path and os.path.isfile(raw_slides_path):
        try:
            with open(raw_slides_path, "r", encoding="utf-8") as f_raw:
                raw_list = json.load(f_raw)
                if isinstance(raw_list, list):
                    raw_db = {s.get("slide_number"): s for s in raw_list if isinstance(s, dict) and "slide_number" in s}
        except Exception:
            pass

    doc = docx.Document()
    
    # Font availability verification
    if not check_font_availability("Dubai"):
        print("⚠️  هشدار فونت: فونت 'Dubai' روی سیستم شما شناسایی نشد. ممکن است نمایش فایل در Word از فونت جایگزین استفاده کند.")
        
    # Auto-detect ref book if not explicitly provided
    if not ref_book:
        corpus = title + " " + " ".join(s.get("title_fa", "") + " " + s.get("title_en", "") for s in slides)
        ref_book = detect_ref_book(corpus)
        print(f"📖 کتاب مرجع ریل ۳ تشخیص داده شد: '{ref_book}'")
        
    add_title(doc, title)
    
    # 1. Build Dynamic Table of Contents (TOC) if section_titles exist
    toc_rows = []
    sec_idx = 1
    for i, s in enumerate(slides):
        if s.get("section_title"):
            stitle = s.get("section_title")
            atime = s.get("audio_time", "")
            snum = s.get("slide_number", i + 1)
            # Find range of slides in this section until next section_title or end
            end_snum = snum
            for j in range(i + 1, len(slides)):
                if slides[j].get("section_title"):
                    break
                end_snum = slides[j].get("slide_number", j + 1)
            slide_range = f"اسلاید {snum}" if snum == end_snum else f"اسلایدهای {snum} الی {end_snum}"
            toc_rows.append((f"بخش {sec_idx}", stitle, atime, slide_range))
            sec_idx += 1
            
    if toc_rows:
        add_toc_table(doc, toc_rows)
        
    # 2. Iterate slides and render Track 1 (Lecture) + Track 2 (Slide Box)
    last_rendered_spoken_norm = None
    for slide in slides:
        num = slide.get("slide_number", 0)
        
        # Section Header (H1)
        if slide.get("section_title"):
            add_h1(doc, slide.get("section_title", ""), audio_time=slide.get("audio_time"))
            last_rendered_spoken_norm = None
            
        # Track 1: Spoken Lecture (with Smart Deduplication & Cluster Awareness)
        spoken_text = slide.get("spoken_lecture")
        parent_ref = slide.get("parent_lecture_slide") or slide.get("cluster_parent")
        
        if spoken_text:
            cleaned_spoken = clean_markdown_text(flatten_spoken_lecture(spoken_text)).strip()
            # Suppress duplicate lecture banner if text is identical to preceding slide or explicitly child of parent slide
            is_dup = (cleaned_spoken and cleaned_spoken == last_rendered_spoken_norm) or (parent_ref is not None)
            if not is_dup:
                add_spoken_lecture(doc, spoken_text, audio_time=slide.get("audio_time"))
                last_rendered_spoken_norm = cleaned_spoken
        elif parent_ref is not None:
            # Clustered continuation under parent slide: nest slide box cleanly without redundant lecture header
            pass
            
        # Track 1: Optional Figure with caption
        if slide.get("track1_figure"):
            fig = slide.get("track1_figure") or {}
            add_figure_with_caption(doc, fig.get("img_path", ""), fig.get("title", ""), fig.get("caption", ""))
            
        # Track 1: Optional Q&A Boxes
        for qa in slide.get("qa_list", []):
            if isinstance(qa, (list, tuple)) and len(qa) >= 2:
                add_qa_box(doc, qa[0], qa[1])
            elif isinstance(qa, dict):
                add_qa_box(doc, qa.get("question", ""), qa.get("answer", ""))
                
        # Track 1: Optional Cross-Reference Callouts (Non-linear lecture progression)
        for cr in slide.get("cross_references", []):
            if isinstance(cr, dict):
                ref_text = cr.get("reason") or cr.get("note", "")
                add_cross_reference_box(doc, cr.get("target_slide", ""), ref_text, direction=cr.get("direction", "forward"))
            elif isinstance(cr, (list, tuple)) and len(cr) >= 2:
                add_cross_reference_box(doc, cr[0], cr[1])
            elif isinstance(cr, str):
                add_cross_reference_box(doc, "", cr)
                
        # Track 2: Slide Box & Metadata Augmentation
        raw_item = raw_db.get(num, {})
        slide_augmented = dict(slide)
        if raw_item:
            if raw_item.get("has_images"):
                slide_augmented["has_images"] = True
            if raw_item.get("has_charts"):
                slide_augmented["has_charts"] = True
            if raw_item.get("has_smartart"):
                slide_augmented["has_smartart"] = True
            if raw_item.get("has_tables"):
                slide_augmented["has_tables"] = True
            if raw_item.get("has_image_table"):
                slide_augmented["has_image_table"] = True
            if raw_item.get("is_pure_visual") or raw_item.get("is_image_only"):
                slide_augmented["is_pure_visual"] = True
            if bool(slide_augmented.get("has_images") or slide_augmented.get("has_charts") or slide_augmented.get("has_smartart") or slide_augmented.get("has_tables") or slide_augmented.get("has_image_table") or slide_augmented.get("is_pure_visual")):
                slide_augmented["is_visual"] = True

        img_path = slide_augmented.get("img_path") or raw_item.get("img_path")
        if img_path and os.path.exists(img_path):
            img_path = convert_image_to_png(img_path)
        elif img_dir:
            img_path = find_and_prepare_slide_image(img_dir, num)
                
        table_data = slide_augmented.get("table_data", None)
        add_slide_box_from_json(doc, slide_augmented, img_path=img_path, table_data=table_data, ref_book_name=ref_book)
        
    doc.save(output_docx_path)
    print(f"✅ جزوه با موفقیت در '{output_docx_path}' ذخیره شد.")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create medical lecture pamphlet docx from JSON.")
    parser.add_argument("--translated", required=True, help="Path to translated_slides.json")
    parser.add_argument("--output", required=True, help="Path to output .docx file")
    parser.add_argument("--title", default="جزوه جامع پزشکی", help="Document Title")
    parser.add_argument("--ref-book", choices=["هاریسون", "برانوالد", "شوارتز", "نلسون", "ویلیامز", "رابینز", "کاتزونگ", "گایتون"], default=None, help="Reference book name")
    parser.add_argument("--img-dir", default="./slide_images", help="Slide images directory")
    parser.add_argument("--raw", default=None, help="Optional path to raw_slides.json for metadata inheritance")
    parser.add_argument("--version", action="version", version=f"%(prog)s {VERSION}")
    
    args = parser.parse_args()
    build_pamphlet_from_json(args.translated, args.output, title=args.title, ref_book=args.ref_book, img_dir=args.img_dir, raw_slides_path=args.raw)
