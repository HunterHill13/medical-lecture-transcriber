#!/usr/bin/env python3
"""
extract_presentation.py: Universal dual-engine presentation extractor for medical study guides.
Supports both PowerPoint (.pptx) and PDF (.pdf) presentations with strict 1-to-1 Page Index Invariance.

Usage:
    uv run --with python-pptx --with pymupdf python extract_presentation.py --input presentation.pdf --output raw_slides.json [--img-dir ./slide_images]
"""

import os
import sys
import json
import argparse
import re

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

_OCR_ENGINE = None

def get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR()
        except Exception:
            _OCR_ENGINE = False
    return _OCR_ENGINE if _OCR_ENGINE is not False else None

def run_ocr_on_image(img_path: str):
    """
    Runs targeted OCR on rendered slide image to extract embedded labels/captions/table text.
    Returns (lines, mean_confidence, is_confident, is_uncertain).
    A solid OCR detection requires mean_confidence >= 0.60 and meaningful alphanumeric tokens >= 2,
    or at least one high-confidence token (>= 0.70).
    Marginal detections (0.35 <= conf < 0.60 or single token) are flagged as uncertain.
    """
    if not img_path or not os.path.exists(img_path):
        return [], 0.0, False, False
    engine = get_ocr_engine()
    if not engine:
        return [], 0.0, False, False
    try:
        results, _ = engine(img_path)
        if not results:
            return [], 0.0, False, False
        lines = []
        confidences = []
        substantive_tokens = []
        for item in results:
            txt = str(item[1]).strip()
            conf = float(item[2])
            if conf >= 0.35 and len(txt) >= 2:
                lines.append(txt)
                confidences.append(conf)
                sub_toks = re.findall(r'[a-zA-Z0-9\u0600-\u06FF]{2,}', txt)
                substantive_tokens.extend(sub_toks)
        if not lines:
            return [], 0.0, False, False
        mean_conf = sum(confidences) / max(1, len(confidences))
        is_confident = (mean_conf >= 0.60 and len(substantive_tokens) >= 2) or any(c >= 0.70 for c in confidences)
        is_uncertain = not is_confident and len(lines) > 0
        return lines, mean_conf, is_confident, is_uncertain
    except Exception:
        return [], 0.0, False, False

def render_pptx_slides_to_images(pptx_path: str, img_dir: str) -> bool:
    """
    Attempts full-slide rendering of PPTX presentations to high-resolution PNGs (slide_01.png, ...)
    via Windows PowerPoint COM or headless LibreOffice.
    Returns True if full-slide images were generated, False otherwise.
    """
    if not img_dir:
        return False
    os.makedirs(img_dir, exist_ok=True)
    abs_pptx = os.path.abspath(pptx_path)
    abs_img_dir = os.path.abspath(img_dir)
    
    # 1. Try Windows PowerPoint COM automation (Native, highest fidelity)
    try:
        import win32com.client
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = None
        try:
            pres = ppt.Presentations.Open(abs_pptx, ReadOnly=True, Untitled=False, WithWindow=False)
            for idx, slide in enumerate(pres.Slides, 1):
                out_path = os.path.join(abs_img_dir, f"slide_{idx:02d}.png")
                slide.Export(out_path, "PNG", 1920, 1080)
            return True
        finally:
            if pres:
                pres.Close()
            ppt.Quit()
    except Exception:
        pass

    # 2. Try headless LibreOffice conversion to PDF, then render with PyMuPDF
    import shutil
    import subprocess
    soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice_bin:
        try:
            import tempfile
            with tempfile.TemporaryDirectory() as tmp_dir:
                cmd = [soffice_bin, "--headless", "--convert-to", "pdf", abs_pptx, "--outdir", tmp_dir]
                res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)
                base = os.path.splitext(os.path.basename(abs_pptx))[0]
                converted_pdf = os.path.join(tmp_dir, f"{base}.pdf")
                if os.path.exists(converted_pdf):
                    import pymupdf
                    doc = pymupdf.open(converted_pdf)
                    for idx, page in enumerate(doc, 1):
                        pix = page.get_pixmap(dpi=150)
                        pix.save(os.path.join(abs_img_dir, f"slide_{idx:02d}.png"))
                    return True
        except Exception:
            pass

    return False

def extract_pptx(pptx_path, img_dir=None):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    
    prs = Presentation(pptx_path)
    slides_data = []
    
    # Attempt full-slide rendering if img_dir provided
    full_rendered = False
    if img_dir:
        full_rendered = render_pptx_slides_to_images(pptx_path, img_dir)
        
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        title = ""
        text_lines = []
        has_tables = False
        has_images = False
        has_charts = False
        has_smartart = False
        image_extraction_error = False
        
        # Check shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        text_lines.append(t)
                        if not title and len(t) < 120 and not t.startswith("•"):
                            title = t
            if shape.has_table:
                has_tables = True
                for row in shape.table.rows:
                    row_txt = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    text_lines.append(" | ".join(row_txt))
            if getattr(shape, "has_chart", False):
                has_charts = True
            if shape.shape_type == getattr(MSO_SHAPE_TYPE, "SMART_ART", 14):
                has_smartart = True
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                has_images = True
                if img_dir and not full_rendered and hasattr(shape, "image"):
                    try:
                        img_bytes = shape.image.blob
                        ext = shape.image.ext or "png"
                        img_filename = f"slide_{slide_num:02d}_img.{ext}"
                        img_path = os.path.join(img_dir, img_filename)
                        if not os.path.exists(img_path):
                            with open(img_path, "wb") as f_img:
                                f_img.write(img_bytes)
                    except Exception as e:
                        sys.stderr.write(f"WARNING: Failed to extract shape image on slide {slide_num}: {e}\n")
                        image_extraction_error = True
                        
        if not title and text_lines:
            title = text_lines[0]
            
        total_chars = sum(len(line) for line in text_lines)
        total_words = sum(len(line.split()) for line in text_lines)
        is_text_empty = (len(text_lines) == 0)
        has_visual_content = bool(has_images or has_charts or has_smartart or has_tables)
        has_visuals = has_visual_content
        is_image_heavy = bool(has_visual_content and total_words < 5)
        
        # Run targeted OCR on visual slides or sparse slides
        ocr_lines = []
        ocr_conf = 0.0
        ocr_confident = False
        ocr_uncertain = False
        candidate_img = None
        if img_dir:
            full_img = os.path.join(img_dir, f"slide_{slide_num:02d}.png")
            if os.path.exists(full_img):
                candidate_img = full_img
            elif is_text_empty or has_images:
                for ext_cand in ("png", "jpg", "jpeg"):
                    single_img = os.path.join(img_dir, f"slide_{slide_num:02d}_img.{ext_cand}")
                    if os.path.exists(single_img):
                        candidate_img = single_img
                        break
        # Run OCR if visual content exists OR if text is sparse (< 5 words)
        if candidate_img and (has_visuals or total_words < 5 or is_text_empty):
            ocr_lines, ocr_conf, ocr_confident, ocr_uncertain = run_ocr_on_image(candidate_img)

        has_image_text = bool(ocr_confident and len(ocr_lines) > 0)
        is_pure_visual = bool(is_text_empty and not has_image_text and has_visual_content and not ocr_uncertain)
        needs_vision = bool(has_visual_content and total_chars < 50) or is_text_empty or has_charts or has_smartart or ocr_uncertain
            
        slide_meta = {
            "slide_number": slide_num,
            "title_raw": title or f"Slide {slide_num}",
            "text_lines": text_lines,
            "ocr_text_lines": ocr_lines,
            "has_images": has_images,
            "full_slide_rendered": full_rendered,
            "has_tables": has_tables,
            "has_image_table": False,
            "has_image_text": has_image_text,
            "is_image_text_bearing": has_image_text,
            "is_pure_visual": is_pure_visual,
            "has_charts": has_charts,
            "has_smartart": has_smartart,
            "is_text_empty": is_text_empty,
            "is_image_only": (is_text_empty and has_visual_content and not has_image_text and not ocr_uncertain),
            "is_image_heavy": is_image_heavy,
            "needs_vision_inspection": needs_vision,
            "ocr_confidence": ocr_conf,
            "ocr_uncertain": ocr_uncertain,
            "needs_student_review": ocr_uncertain,
            "ocr_hint": "⚠️ Visual text detected via OCR (diagram/micrograph labels). Agent MUST translate labels." if has_image_text else ("⚠️ Pure visual slide without text." if is_pure_visual else ""),
            "image_extraction_error": image_extraction_error
        }
        if img_dir and not full_rendered:
            slide_meta["render_diagnostic"] = "FULL_SLIDE_RENDER_UNAVAILABLE"
            
        slides_data.append(slide_meta)
        
    return slides_data

def extract_pdf(pdf_path, img_dir=None):
    import pymupdf
    doc = pymupdf.open(pdf_path)
    slides_data = []
    
    if img_dir:
        os.makedirs(img_dir, exist_ok=True)
        
    for idx, page in enumerate(doc):
        slide_num = idx + 1
        text = page.get_text()
        raw_lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        # Detect images
        image_list = page.get_images(full=True)
        has_images = len(image_list) > 0
        
        # Render high-res page image if requested
        img_path = None
        image_extraction_error = False
        if img_dir:
            try:
                pix = page.get_pixmap(dpi=150)
                img_path = os.path.join(img_dir, f"slide_{slide_num:02d}.png")
                pix.save(img_path)
            except Exception as e:
                sys.stderr.write(f"WARNING: Failed to render page image for slide {slide_num}: {e}\n")
                image_extraction_error = True
            
        # Detect tables (PyMuPDF built-in table finder + text heuristics)
        has_tables = False
        has_image_table = False
        try:
            tabs = page.find_tables()
            if tabs and len(tabs.tables) > 0:
                has_tables = True
        except Exception:
            pass
            
        if not has_tables:
            for l in raw_lines:
                if "\t" in l or " | " in l:
                    has_tables = True
                    break
                    
        total_chars = sum(len(line) for line in raw_lines)
        total_words = sum(len(line.split()) for line in raw_lines)
        has_visual_content = bool(has_images or has_tables)
        has_visuals = has_visual_content
        is_image_heavy = bool(has_visual_content and total_words < 5)
        is_scanned = bool(len(raw_lines) == 0 and has_images)
        
        # Scanned table vs Native vector table:
        # has_image_table is True ONLY if the table is in a scanned page or lacks native text lines
        if has_tables and (len(raw_lines) == 0 or is_scanned):
            has_image_table = True

        # Run targeted OCR evidence extraction if visual content exists OR text is sparse (< 5 words)
        ocr_lines = []
        ocr_conf = 0.0
        ocr_confident = False
        ocr_uncertain = False
        temp_img_created = False
        ocr_img_path = img_path
        if has_visuals or total_words < 5 or len(raw_lines) == 0:
            if not ocr_img_path:
                import tempfile
                pix_tmp = page.get_pixmap(dpi=150)
                tmp_fd, ocr_img_path = tempfile.mkstemp(suffix=f"_slide_{slide_num:02d}.png")
                os.close(tmp_fd)
                pix_tmp.save(ocr_img_path)
                temp_img_created = True

            if ocr_img_path and os.path.exists(ocr_img_path):
                ocr_lines, ocr_conf, ocr_confident, ocr_uncertain = run_ocr_on_image(ocr_img_path)
                if temp_img_created and os.path.exists(ocr_img_path):
                    try:
                        os.remove(ocr_img_path)
                    except Exception:
                        pass

        has_image_text = bool(ocr_confident and len(ocr_lines) > 0)
        is_pure_visual = bool(len(raw_lines) == 0 and not has_image_text and has_visual_content and not ocr_uncertain)
        needs_vision = bool(has_visual_content and total_chars < 50) or is_scanned or has_image_text or ocr_uncertain
        
        title = raw_lines[0] if raw_lines else (ocr_lines[0] if ocr_lines else f"Slide {slide_num}")
        # Filter out plain numbers or headers
        if len(raw_lines) > 1 and (len(title) <= 2 or title.isdigit()):
            title = raw_lines[1]
        elif len(raw_lines) == 0 and len(ocr_lines) > 1 and (len(title) <= 2 or title.isdigit()):
            title = ocr_lines[1]
                
        slides_data.append({
            "slide_number": slide_num,
            "title_raw": title,
            "text_lines": raw_lines,
            "ocr_text_lines": ocr_lines,
            "has_images": has_images,
            "has_tables": has_tables,
            "has_image_table": has_image_table,
            "has_image_text": has_image_text,
            "is_image_text_bearing": has_image_text,
            "is_pure_visual": is_pure_visual,
            "is_image_only": (len(raw_lines) == 0 and has_visual_content and not has_image_text and not ocr_uncertain),
            "is_image_heavy": is_image_heavy,
            "is_scanned": is_scanned,
            "needs_vision_inspection": needs_vision,
            "ocr_confidence": ocr_conf,
            "ocr_uncertain": ocr_uncertain,
            "needs_student_review": ocr_uncertain,
            "student_review_note": "⚠️ بازبینی دانشجو: متن استخراج‌شده با OCR دارای عدم قطعیت آماری یا کیفیت پایین است." if ocr_uncertain else "",
            "ocr_hint": "⚠️ Visual text detected via OCR (diagram/micrograph labels). Agent MUST translate labels." if has_image_text else ("⚠️ Scanned/image page without text layer." if is_scanned else ""),
            "image_extraction_error": image_extraction_error
        })
        
    return slides_data

def main():
    parser = argparse.ArgumentParser(description="Universal dual-engine presentation extractor")
    parser.add_argument("--input", "-i", required=True, help="Input PPTX or PDF file")
    parser.add_argument("--output", "-o", default="raw_slides.json", help="Output JSON path")
    parser.add_argument("--img-dir", help="Directory to save extracted slide images")
    args = parser.parse_args()
    
    if not os.path.exists(args.input):
        print(f"Error: Input file '{args.input}' not found!", file=sys.stderr)
        sys.exit(1)
        
    ext = os.path.splitext(args.input)[1].lower()
    if ext == ".pptx":
        print(f"Extracting PPTX presentation: {args.input}")
        data = extract_pptx(args.input, args.img_dir)
    elif ext == ".pdf":
        print(f"Extracting PDF presentation: {args.input}")
        data = extract_pdf(args.input, args.img_dir)
    else:
        print(f"Error: Unsupported file format '{ext}'. Must be .pptx or .pdf", file=sys.stderr)
        sys.exit(1)
        
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        
    print(f"Successfully extracted {len(data)} slides into '{args.output}'.")
    image_only_count = sum(1 for s in data if s.get("is_image_only", False))
    print(f"Total slides: {len(data)} | Image-only slides: {image_only_count}")

if __name__ == "__main__":
    main()
