#!/usr/bin/env python3
"""
annotate_docx.py: Injects structured, beautifully styled RTL callout boxes
into a Word document containing PowerPoint slide texts, without modifying
or shifting any of the original slide paragraphs, tables, or drawings.
"""

import os
import sys
import json
import re
import html
import argparse
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from text_utils import normalize_persian_text

def style_p(p, space_before=2, space_after=2, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    # Remove existing bidi tags to avoid duplicate XML nodes if called multiple times
    existing = pPr.find(qn('w:bidi'))
    if existing is not None:
        pPr.remove(existing)
    pPr.append(parse_xml(r'<w:bidi {} w:val="1"/>'.format(nsdecls('w'))))

def add_run(p, text, font_name="Dubai", size_pt=10.5, bold=False, italic=False, color_rgb=(0x26, 0x26, 0x26)):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml(r'<w:rtl {} w:val="1"/>'.format(nsdecls('w'))))
    safe_font_name = html.escape(str(font_name), quote=True)
    rFonts = parse_xml(r'<w:rFonts {} w:ascii="{}" w:hAnsi="{}" w:cs="{}"/>'.format(nsdecls('w'), safe_font_name, safe_font_name, safe_font_name))
    rPr.append(rFonts)
    sz_half_pts = str(int(size_pt * 2))
    rPr.append(parse_xml(r'<w:szCs {} w:val="{}"/>'.format(nsdecls('w'), sz_half_pts)))
    if bold:
        rPr.append(parse_xml(r'<w:bCs {}/>'.format(nsdecls('w'))))
    return run

def create_callout_cell(doc, target_p=None, border_color="1F4E79", bg_color="F7F9FB"):
    tbl = doc.add_table(rows=1, cols=1)
    if target_p is not None:
        target_p._p.addprevious(tbl._tbl)
    
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(r'<w:tblW {} w:w="5000" w:type="pct"/>'.format(nsdecls('w'))))
    tblPr.append(parse_xml(r'<w:bidiVisual {}/>'.format(nsdecls('w'))))
    
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color)))
    tcPr.append(parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="single" w:sz="36" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(nsdecls('w'), border_color)))
    tcPr.append(parse_xml(r'''
        <w:tcMar {} >
            <w:top w:w="160" w:type="dxa"/>
            <w:bottom w:w="160" w:type="dxa"/>
            <w:left w:w="220" w:type="dxa"/>
            <w:right w:w="260" w:type="dxa"/>
        </w:tcMar>
    '''.format(nsdecls('w'))))
    return cell

def render_box(cell, box_data):
    # Header
    p0 = cell.paragraphs[0]
    style_p(p0, space_before=4, space_after=4)
    add_run(p0, box_data["header"], font_name="Calibri", size_pt=12, bold=True, color_rgb=(0x1F, 0x4E, 0x79))
    
    # Sections
    for sec in box_data.get("sections", []):
        if "title" in sec and sec["title"]:
            sp = cell.add_paragraph()
            style_p(sp, space_before=4, space_after=2)
            title_color = tuple(sec.get("title_color", [14, 98, 81]))
            add_run(sp, sec["title"], font_name="Calibri", size_pt=11, bold=True, color_rgb=title_color)
        
        for item in sec.get("items", []):
            ip = cell.add_paragraph()
            style_p(ip, space_before=1.5, space_after=2)
            if isinstance(item, list) and len(item) == 2:
                lead, body = item
                lead_color = tuple(sec.get("lead_color", [26, 82, 118]))
                add_run(ip, lead + " ", font_name="Calibri", size_pt=10.5, bold=True, color_rgb=lead_color)
                add_run(ip, body, font_name="Calibri", size_pt=10.5, bold=False, color_rgb=(0x26, 0x26, 0x26))
            else:
                add_run(ip, str(item), font_name="Calibri", size_pt=10.5, bold=False, color_rgb=(0x26, 0x26, 0x26))

def extract_page_number(text):
    """Extracts slide or page number from English or Persian text:
    Matches 'Page_1', 'Page 1', 'اسلاید 1', 'اسلاید مرتبط ۱', 'صفحه ۱', etc."""
    if not text:
        return None
    cleaned = normalize_persian_text(str(text))
    
    # 1. Match Page_X or Page X
    m = re.search(r'\bPage[_\s]*(\d+)\b', cleaned, re.IGNORECASE)
    if m:
        return int(m.group(1))
    
    # 2. Match Persian: اسلاید مرتبط X or اسلاید X or صفحه X
    m = re.search(r'(?:اسلاید(?:[\s_]*مرتبط)?|صفحه)[\s_]*(\d+)', cleaned)
    if m:
        return int(m.group(1))
        
    return None

def annotate_docx(input_docx, annotations_json, output_docx):
    """
    Annotates Word document with structured RTL callout boxes before matching markers.
    
    Note: Currently scans document-level paragraphs (doc.paragraphs). Markers embedded
    inside existing Word tables or nested cells are not directly matched.
    """
    if not os.path.isfile(input_docx):
        raise FileNotFoundError(f"Input docx not found: {input_docx}")
    if not os.path.isfile(annotations_json):
        raise FileNotFoundError(f"Annotations JSON not found: {annotations_json}")

    doc = docx.Document(input_docx)
    with open(annotations_json, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    all_paras = [(p.text.strip(), p) for p in doc.paragraphs]
    page_paras = []
    for p in doc.paragraphs:
        t = p.text.strip()
        p_num = extract_page_number(t)
        if p_num is not None:
            page_paras.append((p_num, t, p))
            
    boxes = data.get("boxes", [])
    end_box = data.get("end_box", None)
    matched_boxes = 0
    
    for bdata in boxes:
        marker = bdata["before_marker"]
        marker_num = extract_page_number(marker)
        
        target_p = None
        if marker_num is not None:
            # Deterministic numeric matching (prevents Page_1 matching Page_10, avoids None==None)
            for p_num, p_text, p in page_paras:
                if p_num == marker_num:
                    target_p = p
                    break
        else:
            # Exact or word-boundary fallback for non-page markers across ALL document paragraphs
            marker_clean = marker.strip()
            for p_text, p in all_paras:
                if marker_clean == p_text:
                    target_p = p
                    break
            if not target_p:
                for p_text, p in all_paras:
                    if re.search(r'\b' + re.escape(marker_clean) + r'\b', p_text, re.IGNORECASE):
                        target_p = p
                        break
                        
        if target_p:
            cell = create_callout_cell(doc, target_p=target_p)
            render_box(cell, bdata)
            matched_boxes += 1
        else:
            print(f"Warning: Marker '{marker}' (detected page #{marker_num}) not found in docx.")
                
    if end_box:
        end_cell = create_callout_cell(doc, target_p=None)
        render_box(end_cell, end_box)
        
    if len(boxes) > 0 and matched_boxes == 0 and not end_box:
        raise RuntimeError(f"Annotation failed: None of the {len(boxes)} annotation boxes matched any paragraphs in '{input_docx}'.")

    doc.save(output_docx)
    print(f"Success! Document annotated ({matched_boxes}/{len(boxes)} boxes matched) and saved to: {output_docx}")

def main():
    parser = argparse.ArgumentParser(description="Annotate slide docx with structured RTL callout notes.")
    parser.add_argument("input_docx", help="Path to original docx file")
    parser.add_argument("annotations_json", help="Path to JSON file with annotations")
    parser.add_argument("--output-docx", required=True, help="Path to save annotated docx file")
    
    args = parser.parse_args()
    try:
        annotate_docx(args.input_docx, args.annotations_json, args.output_docx)
    except Exception as e:
        print(f"❌ Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
